"""Forfait en équipes (3,8 par défaut) : est-ce qu'on dépasse avec la majoration des tickets ?

Lit une ou plusieurs extractions Praxedo (une par mois, .csv séparé par des
points-virgules ou .xlsx) et produit un classeur Excel avec formules :

    python3 forfait-equipes/calcul_forfait.py export_JUIN.csv export_SEPTEMBRE.csv
    python3 forfait-equipes/calcul_forfait.py *.csv -o Forfait_equipes.xlsx
    python3 forfait-equipes/calcul_forfait.py *.csv --reprendre ancien_classeur_rempli.xlsx
    python3 forfait-equipes/calcul_forfait.py glaucos/*.csv --nom Glaucos --forfait 0.05
    python3 forfait-equipes/calcul_forfait.py losange/*.csv --nom Losange --reel --cerfa-devis --prix PRIX_EQUIPE_MOIS

Règles appliquées (données par le responsable) :
- Avec --cerfa-devis (ex. Losange), les tickets marqués « Oui » dans la
  colonne Cerfa / Plainte / Constat sont payés sur devis : listés, mais
  exclus du calcul comme les IRIS.
- Les tickets dont le numéro commence par MAL- sont des IRIS, payés à part :
  ils sont listés mais ne comptent pas dans les équipes.
- Pour chaque équipe et chaque jour, on prend le temps réellement occupé par
  ses tickets (de Début à Soldée, les chevauchements comptés une seule fois).
  Moins de 2 h 30 = 1 ticket, puis +1 ticket à 2 h 30, 4 h 30, 6 h 30, etc.
  Le nombre de tickets comptés ne descend jamais sous le nombre de tickets
  faits : la bonification de la journée est la différence.
- Si le dernier ticket de la journée est soldé après 16 h, l'équipe a
  dépassé sa journée : ses tickets de 2 h 30 et plus sont alors majorés
  même si elle a fait d'autres tickets (leur temps est regroupé : deux
  tickets de 2 h 30 = un seul majoré).
- Un ticket soldé le lendemain avant 6 h est du travail de nuit : son temps
  compte en entier sur le jour où il a commencé.
- Un autre ticket ouvert sur plusieurs jours compte 1 ticket sans
  majoration, et part dans l'onglet « À vérifier » où l'on peut saisir le
  temps réellement travaillé.
- Équipes du mois = tickets ÷ 4 ÷ jours ouvrés (lundi à vendredi hors jours
  fériés). Les tickets des week-ends et jours fériés comptent quand même.
"""

import argparse
import collections
import csv
import datetime as dt
import os
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FORFAIT = 3.8
NOM = ""  # nom du contrat, affiché dans les titres (ex. « Glaucos »)
PRIX = None  # payé au réel : prix d'une équipe par mois, en euros
TICKETS_PAR_JOUR = 4
SEUIL_H = 2.5
TRANCHE_H = 2
PREFIXE_IRIS = "MAL-"
CERFA_DEVIS = False  # --cerfa-devis : tickets « Cerfa / Plainte / Constat = Oui » payés sur devis, exclus
# Un ticket soldé le lendemain avant cette heure est du travail de nuit
# (ex. 16 h - 0 h 25) : son temps compte en entier sur le jour où il a commencé.
FIN_NUIT = dt.time(6, 0)
FIN_JOURNEE = dt.time(16, 0)
# Noms Praxedo qui regroupent plusieurs équipes travaillant en même temps,
# quand l'extraction n'a pas la colonne « Prénom technicien ».
# Leurs tickets sont répartis en équipes A, B… d'après l'enchaînement des
# horaires (A = celle qui commence en premier dans la journée).
EQUIPES_MULTIPLES = {"[RES] SADOK": 2}

MOIS = ["Janvier", "Février", "Mars", "Avril", "Mai", "Juin", "Juillet",
        "Août", "Septembre", "Octobre", "Novembre", "Décembre"]
JOURS = ["lun.", "mar.", "mer.", "jeu.", "ven.", "sam.", "dim."]

P = "'Paramètres'"
D = "'Détail tickets'"
EJ = "'Par équipe et jour'"
AV = "'À vérifier'"


# --- Lecture des extractions ---------------------------------------------

def lire_date(v):
    if isinstance(v, dt.datetime):
        return v
    v = str(v or "").strip()
    for fmt in ("%d/%m/%Y %H:%M", "%d/%m/%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S"):
        try:
            return dt.datetime.strptime(v, fmt)
        except ValueError:
            pass
    return None


def lire_lignes(chemin):
    if chemin.lower().endswith((".xlsx", ".xlsm")):
        ws = load_workbook(chemin, read_only=True, data_only=True).active
        it = ws.iter_rows(values_only=True)
        entetes = [str(h or "").strip() for h in next(it)]
        return [dict(zip(entetes, ligne)) for ligne in it]
    with open(chemin, encoding="utf-8-sig", newline="") as f:
        lignes = list(csv.DictReader(f, delimiter=";"))
    return [{(k or "").strip(): v for k, v in l.items()} for l in lignes]


def lire_export(chemin):
    tickets, ignores = [], []
    for l in lire_lignes(chemin):
        num = str(l.get("N°") or "").strip()
        if not num:
            continue
        t = {
            "num": num,
            "equipe": str(l.get("Nom technicien") or "").strip() or "(sans nom)",
            "prenom": str(l.get("Prénom technicien") or "").strip(),
            "site": str(l.get("Site") or "").strip(),
            "type": str(l.get("Type inter.") or "").strip(),
            "statut": str(l.get("Statut") or "").strip(),
            "cerfa": str(l.get("Cerfa / Plainte/ Constat") or "").strip().lower() == "oui",
            "debut": lire_date(l.get("Début")),
            "fin": lire_date(l.get("Soldée")),
        }
        if "annul" in t["statut"].lower() or not t["debut"] or not t["fin"] or t["fin"] < t["debut"]:
            ignores.append(t)
            continue
        t["iris"] = num.upper().startswith(PREFIXE_IRIS)
        t["devis"] = CERFA_DEVIS and t["cerfa"] and not t["iris"]
        t["exclu"] = t["iris"] or t["devis"]
        lendemain = t["fin"].date() - t["debut"].date() == dt.timedelta(1)
        t["nuit"] = lendemain and t["fin"].time() <= FIN_NUIT
        t["multi"] = t["debut"].date() != t["fin"].date() and not t["nuit"]
        tickets.append(t)
    if not tickets:
        sys.exit(f"Aucun ticket lisible dans {chemin}")
    annee, mois = collections.Counter((t["fin"].year, t["fin"].month) for t in tickets).most_common(1)[0][0]
    return (annee, mois), tickets, ignores


# --- Calculs ---------------------------------------------------------------

def paques(a):
    g, c = a % 19, a // 100
    h = (c - c // 4 - (8 * c + 13) // 25 + 19 * g + 15) % 30
    i = h - (h // 28) * (1 - (h // 28) * (29 // (h + 1)) * ((21 - g) // 11))
    j = (a + a // 4 + i + 2 - c + c // 4) % 7
    m = 3 + (i - j + 40) // 44
    return dt.date(a, m, i - j + 28 - 31 * (m // 4))


def jours_feries(annee):
    p = paques(annee)
    return [
        (dt.date(annee, 1, 1), "Jour de l'an"),
        (p + dt.timedelta(1), "Lundi de Pâques"),
        (dt.date(annee, 5, 1), "Fête du travail"),
        (dt.date(annee, 5, 8), "Victoire 1945"),
        (p + dt.timedelta(39), "Ascension"),
        (p + dt.timedelta(50), "Lundi de Pentecôte"),
        (dt.date(annee, 7, 14), "Fête nationale"),
        (dt.date(annee, 8, 15), "Assomption"),
        (dt.date(annee, 11, 1), "Toussaint"),
        (dt.date(annee, 11, 11), "Armistice"),
        (dt.date(annee, 12, 25), "Noël"),
    ]


def temps_occupe(plages):
    """Heures couvertes par les plages, chevauchements comptés une fois."""
    total, debut, fin = dt.timedelta(), None, None
    for a, b in sorted(plages):
        if debut is None:
            debut, fin = a, b
        elif a <= fin:
            fin = max(fin, b)
        else:
            total += fin - debut
            debut, fin = a, b
    if debut is not None:
        total += fin - debut
    return round(total.total_seconds() / 3600, 4)


def nommer_par_prenom(tickets):
    """Ajoute le prénom au nom quand un même nom couvre plusieurs prénoms
    dans le mois (ex. « [RES] SADOK (Mohamed) » et « [RES] SADOK (Kada) »)."""
    prenoms = collections.defaultdict(set)
    for t in tickets:
        if t["prenom"]:
            prenoms[t["equipe"]].add(t["prenom"])
    for t in tickets:
        if len(prenoms[t["equipe"]]) > 1 and t["prenom"]:
            t["equipe"] = f"{t['equipe']} ({t['prenom']})"


def separer_equipes(tickets):
    """Répartit les tickets des noms de EQUIPES_MULTIPLES entre leurs équipes,
    quand l'extraction n'a pas de prénom pour les distinguer.

    Jour par jour, un ticket rejoint l'équipe qui vient de finir son ticket
    précédent (5 min de marge), sinon une nouvelle équipe tant qu'il en reste,
    sinon celle déjà sur le même site, sinon celle qui finit le plus tôt.
    """
    marge = dt.timedelta(minutes=5)
    jours = collections.defaultdict(list)
    for t in tickets:
        if t["equipe"] in EQUIPES_MULTIPLES:
            jours[(t["equipe"], t["debut"].date())].append(t)
    for (nom, _), ts in jours.items():
        pistes = []  # [fin du dernier ticket, site du dernier ticket]
        for t in sorted(ts, key=lambda t: (t["multi"], t["debut"], t["fin"])):
            fin = t["debut"] if t["multi"] else t["fin"]
            meme_site = [i for i, (f, s) in enumerate(pistes) if s == t["site"] and f > t["debut"] - marge]
            libres = [i for i, (f, _) in enumerate(pistes) if f <= t["debut"] + marge]
            if libres:
                i = max(libres, key=lambda i: pistes[i][0])
            elif len(pistes) < EQUIPES_MULTIPLES[nom]:
                pistes.append([fin, t["site"]])
                i = len(pistes) - 1
            elif meme_site:
                i = meme_site[0]
            else:
                i = min(range(len(pistes)), key=lambda i: pistes[i][0])
            pistes[i] = [max(pistes[i][0], fin), t["site"]]
            t["equipe"] = f"{nom} {'ABCDEFGH'[i]}"


def tickets(heures):
    """Tickets pour une durée : 1 sous le seuil, puis +1 par tranche."""
    h = round(heures, 4)
    return 1 if h < SEUIL_H else 2 + int((h - SEUIL_H) // TRANCHE_H)


def duree(t):
    return (t["fin"] - t["debut"]).total_seconds() / 3600


def longs(g):
    return [t for t in g if round(duree(t), 4) >= SEUIL_H]


def tickets_journee(g):
    """Tickets comptés pour une équipe et un jour (même règle que le classeur)."""
    compte = max(len(g), tickets(temps_occupe([(t["debut"], t["fin"]) for t in g])))
    jour = min(t["debut"] for t in g).date()
    if max(t["fin"] for t in g) > dt.datetime.combine(jour, FIN_JOURNEE) and longs(g):
        lg = longs(g)
        compte = max(compte, len(g) - len(lg) + tickets(temps_occupe([(t["debut"], t["fin"]) for t in lg])))
    return compte


def libelle_mois(cle):
    return f"{MOIS[cle[1] - 1]} {cle[0]}"


# --- Classeur ----------------------------------------------------------------

POLICE = "Arial"
F_NORMAL = Font(name=POLICE, size=10)
F_GRAS = Font(name=POLICE, size=10, bold=True)
F_TITRE = Font(name=POLICE, size=14, bold=True)
F_NOTE = Font(name=POLICE, size=9, italic=True, color="555555")
F_SAISIE = Font(name=POLICE, size=10, color="0000FF")
F_LIEN = Font(name=POLICE, size=10, color="008000")
FOND_ENTETE = PatternFill("solid", fgColor="DDE4EE")
FOND_SAISIE = PatternFill("solid", fgColor="FFFF00")
FOND_TOTAL = PatternFill("solid", fgColor="F2F2F2")
TRAIT = Side(style="thin", color="BBBBBB")
BORD = Border(top=TRAIT, bottom=TRAIT, left=TRAIT, right=TRAIT)
F_DATE = "dd/mm/yyyy"
F_DATEH = "dd/mm/yyyy hh:mm"
F_HEURE = "hh:mm"
F_DEC = "0.00"
F_EUR = '#,##0 "€"'
F_EQ = F_DEC  # format des équipes : 3 décimales pour un petit forfait (voir main)


def nb_dec():
    return 3 if FORFAIT is not None and FORFAIT < 1 else 2


def titre_forfait():
    if FORFAIT is None:
        return f"{NOM + ' : ' if NOM else ''}payé au réel"
    return f"{NOM + ' : ' if NOM else ''}forfait {FORFAIT:g} équipes".replace(".", ",")


def entetes(ws, ligne, noms, largeurs):
    for i, (nom, larg) in enumerate(zip(noms, largeurs), 1):
        c = ws.cell(ligne, i, nom)
        c.font, c.fill, c.border = F_GRAS, FOND_ENTETE, BORD
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = larg
    ws.row_dimensions[ligne].height = 42
    ws.freeze_panes = ws.cell(ligne + 1, 1)


def ecrire(ws, ligne, valeurs, formats=None, police=F_NORMAL):
    for i, v in enumerate(valeurs, 1):
        c = ws.cell(ligne, i, v)
        c.font, c.border = police, BORD
        if formats and formats.get(i):
            c.number_format = formats[i]


def formule_tickets(cellule):
    """Tickets comptés pour une durée : 1 sous le seuil, puis +1 par tranche."""
    h = f"ROUND({cellule},4)"
    return f"IF({h}<{P}!$C$5,1,2+INT(({h}-{P}!$C$5)/{P}!$C$6))"


def graphique(syn, fin_syn, ligne, dest=None, colonne="A"):
    """Équipes par mois, sans et avec majoration, et la ligne du forfait."""
    mois = Reference(syn, min_col=1, min_row=5, max_row=fin_syn)
    barres = BarChart()
    barres.type, barres.grouping, barres.overlap, barres.gapWidth = "col", "clustered", -10, 60
    barres.title = ("Équipes par mois : sans et avec majoration" if FORFAIT is None
                    else "Équipes par mois : sans et avec majoration, face au forfait")
    barres.y_axis.title = "Équipes"
    barres.y_axis.scaling.min = 0
    barres.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="E6E6E3"))
    for col, couleur in ((6, "6DA7EC"), (9, "1C5CAB")):
        barres.add_data(Reference(syn, min_col=col, min_row=4, max_row=fin_syn), titles_from_data=True)
        s = barres.series[-1]
        s.graphicalProperties.solidFill = couleur
        s.graphicalProperties.line.noFill = True
    barres.set_categories(mois)
    if FORFAIT is not None:
        forfait = LineChart()
        forfait.add_data(Reference(syn, min_col=10, min_row=4, max_row=fin_syn), titles_from_data=True)
        s = forfait.series[0]
        s.graphicalProperties.line.solidFill = "52514E"
        s.graphicalProperties.line.width = 22000
        s.graphicalProperties.line.prstDash = "dash"
        s.smooth = False
        forfait.y_axis.axId = barres.y_axis.axId
        barres += forfait
    barres.x_axis.delete = barres.y_axis.delete = False
    barres.legend.position = "b"
    barres.width, barres.height = 24, 11
    (dest or syn).add_chart(barres, f"{colonne}{ligne}")


def graphique_tickets(syn, fin_syn, dest, ancre):
    """Tickets faits + bonifications par mois, face aux tickets couverts par le forfait."""
    barres = BarChart()
    barres.type, barres.grouping, barres.overlap, barres.gapWidth = "col", "stacked", 100, 60
    barres.title = ("Tickets par mois : faits + bonifications" if FORFAIT is None
                    else "Tickets par mois : faits + bonifications, face au forfait")
    barres.y_axis.title = "Tickets"
    barres.y_axis.scaling.min = 0
    barres.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="E6E6E3"))
    for col, couleur in ((5, "6DA7EC"), (7, "1C5CAB")):
        barres.add_data(Reference(syn, min_col=col, min_row=4, max_row=fin_syn), titles_from_data=True)
        s = barres.series[-1]
        s.graphicalProperties.solidFill = couleur
        s.graphicalProperties.line.solidFill = "FFFFFF"
    barres.set_categories(Reference(syn, min_col=1, min_row=5, max_row=fin_syn))
    if FORFAIT is not None:
        forfait = LineChart()
        forfait.add_data(Reference(syn, min_col=11, min_row=4, max_row=fin_syn), titles_from_data=True)
        s = forfait.series[0]
        s.graphicalProperties.line.solidFill = "52514E"
        s.graphicalProperties.line.width = 22000
        s.graphicalProperties.line.prstDash = "dash"
        s.smooth = False
        forfait.y_axis.axId = barres.y_axis.axId
        barres += forfait
    barres.x_axis.delete = barres.y_axis.delete = False
    barres.legend.position = "b"
    barres.width, barres.height = 24, 11
    dest.add_chart(barres, ancre)


def tableau_de_bord(wb, mois, fin_syn, fin_det, fin_ej):
    """Onglet d'indicateurs : chiffres clés, évolution par mois, par équipe."""
    tb = wb.create_sheet("Tableau de bord", 0)
    S, tot = "'Synthèse'", fin_syn + 1
    plage = lambda col: f"{S}!${col}$5:${col}${fin_syn}"
    tb.sheet_view.showGridLines = False
    tb.column_dimensions["A"].width = 2
    for i in range(2, 15):
        tb.column_dimensions[get_column_letter(i)].width = 12.5
    tb.column_dimensions["B"].width = 16
    tb["B1"] = f"Tableau de bord {titre_forfait()}"
    tb["B1"].font = Font(name=POLICE, size=16, bold=True)
    tb["B2"] = f'="Période : "&{S}!A5&" à "&{S}!A{fin_syn}&" · "&COUNTA({plage("A")})&" mois · règles dans l\'onglet Synthèse"'
    tb["B2"].font = F_NOTE

    # Chiffres clés : 6 cases de 2 colonnes.
    fond = PatternFill("solid", fgColor="F3F6FA")
    cases_reel = [
        ("Équipes facturables en moyenne (avec majoration)", f"={S}!I{tot}", F_EQ,
         f'="sans majoration : "&FIXED({S}!F{tot},{nb_dec()})'),
        ("Montant facturable sur la période", f"={S}!L{tot}", F_EUR,
         f'=IF({P}!$C$9="","prix d\'une équipe à saisir dans Paramètres","sans majoration : "&FIXED({S}!K{tot},0)&" €")'),
        ("Gain de la majoration", f"={S}!M{tot}", F_EUR,
         f'="+"&FIXED({S}!J{tot},{nb_dec()})&" équipe en moyenne"'),
        ("Mois le plus chargé", f'=INDEX({plage("A")},MATCH(MAX({plage("I")}),{plage("I")},0))', None,
         f'=FIXED(MAX({plage("I")}),{nb_dec()})&" équipes avec majoration"'),
        ("Apport de la majoration (tickets)", f'="+"&{S}!G{tot}&" tickets"', None,
         f'="soit +"&FIXED({S}!G{tot}/{S}!E{tot}*100,0)&" % des tickets faits"'),
        ("Dernier mois", f"={S}!I{fin_syn}", F_EQ,
         f'={S}!A{fin_syn}&" · "&IF({S}!I{fin_syn}>={S}!I{max(fin_syn - 1, 5)},"+","")'
         f'&FIXED({S}!I{fin_syn}-{S}!I{max(fin_syn - 1, 5)},{nb_dec()})&" vs mois précédent"'),
    ]
    cases = [
        ("Équipes en moyenne, avec majoration", f"={S}!I{tot}", F_EQ,
         f'="sans majoration : "&FIXED({S}!F{tot},{nb_dec()})&" · forfait "&FIXED({P}!$C$3,2)'),
        ("Mois au-dessus du forfait", f'=COUNTIF({plage("I")},">"&{P}!$C$3)&" sur "&COUNTA({plage("A")})', None,
         f'="seuil : "&FIXED({P}!$C$3,2)&" équipes"'),
        ("Mois le plus chargé", f'=INDEX({plage("A")},MATCH(MAX({plage("I")}),{plage("I")},0))', None,
         f'=FIXED(MAX({plage("I")}),{nb_dec()})&" équipes avec majoration"'),
        ("Apport de la majoration", f'="+"&FIXED({S}!I{tot}-{S}!F{tot},{nb_dec()})&" équipe"', None,
         f'="+"&{S}!G{tot}&" tickets, soit +"&FIXED({S}!G{tot}/{S}!E{tot}*100,0)&" %"'),
        ("Dernier mois", f"={S}!I{fin_syn}", F_EQ,
         f'={S}!A{fin_syn}&" · "&IF({S}!I{fin_syn}>={S}!I{max(fin_syn - 1, 5)},"+","")'
         f'&FIXED({S}!I{fin_syn}-{S}!I{max(fin_syn - 1, 5)},{nb_dec()})&" vs mois précédent"'),
        ("Marge sur la période (tickets)", f"=-{S}!L{tot}", "0;-0;0",
         f'=IF({S}!L{tot}<0,"encore sous le forfait","au-dessus du forfait")'),
    ]
    for i, (titre, valeur, fmt, sous) in enumerate(cases_reel if FORFAIT is None else cases):
        c1, c2 = get_column_letter(2 + 2 * i), get_column_letter(3 + 2 * i)
        for r, contenu, police in ((4, titre, Font(name=POLICE, size=9, color="52514E")),
                                   (5, valeur, Font(name=POLICE, size=20 if fmt else 15, bold=True)),
                                   (6, sous, Font(name=POLICE, size=9, color="52514E"))):
            tb.merge_cells(f"{c1}{r}:{c2}{r}")
            c = tb[f"{c1}{r}"]
            c.value, c.font = contenu, police
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if fmt and r == 5:
                c.number_format = fmt
            for col in (c1, c2):
                tb[f"{col}{r}"].fill = fond
    tb.row_dimensions[4].height, tb.row_dimensions[5].height, tb.row_dimensions[6].height = 26, 34, 26
    rouge, vert = PatternFill("solid", fgColor="F8D7DA"), PatternFill("solid", fgColor="D4EDDA")
    for zone, formule in () if FORFAIT is None else (("B5:C5", f"$B$5>{P}!$C$3"), ("J5:K5", f"$J$5>{P}!$C$3")):
        tb.conditional_formatting.add(zone, FormulaRule(formula=[formule], fill=rouge))
        tb.conditional_formatting.add(zone, FormulaRule(formula=[f"NOT({formule})"], fill=vert))

    # Évolution mois par mois.
    tb["B8"] = "Évolution mois par mois"
    tb["B8"].font = Font(name=POLICE, size=12, bold=True)
    noms = ["Mois", "Tickets faits", "Bonifications", "Équipes sans majoration", "Équipes avec majoration",
            "Évolution vs mois précédent", "Montant avec majoration (€)" if FORFAIT is None else "Part du forfait utilisée",
            "Journées finies après 16 h", "Tickets de 2 h 30 et plus", "Tickets par jour ouvré", "Tickets à vérifier",
            "Gain de la majoration (€)" if FORFAIT is None else "Résultat"]
    reel = FORFAIT is None
    f_h = (lambda r, k: f"={S}!L{k}") if reel else (lambda r, k: f"=F{r}/{P}!$C$3")
    f_res = {8: F_EUR, 13: F_EUR} if reel else {8: "0%"}
    for i, nom in enumerate(noms):
        c = tb.cell(9, 2 + i, nom)
        c.font, c.fill, c.border = F_GRAS, FOND_ENTETE, BORD
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    tb.row_dimensions[9].height = 42
    tb.column_dimensions["M"].width = 22
    eA, eJ = f"{EJ}!$A$4:$A${fin_ej}", f"{EJ}!$J$4:$J${fin_ej}"
    dA, dJ, dI = (f"{D}!${c}$2:${c}${fin_det}" for c in "AJI")
    r = 10
    for k in range(5, fin_syn + 1):
        ecrire(tb, r, [None, f"={S}!A{k}", f"={S}!E{k}", f"={S}!G{k}", f"={S}!F{k}", f"={S}!I{k}",
                       "–" if r == 10 else f"=F{r}-F{r - 1}", f_h(r, k),
                       f"=IFERROR(COUNTIFS({eA},B{r},{eJ},\"Oui\")/COUNTIFS({eA},B{r}),0)",
                       f'=IFERROR(COUNTIFS({dA},B{r},{dJ},"Non",{dI},">="&{P}!$C$5)/C{r},0)',
                       f"=C{r}/{S}!C{k}", f"={S}!N{k}", f"={S}!M{k}"],
               {5: F_EQ, 6: F_EQ, 7: f"+{F_EQ};-{F_EQ};{F_EQ}", 9: "0%", 10: "0%", 11: "0.0", **f_res})
        tb.cell(r, 1).border = Border()
        r += 1
    fin_tb = r - 1
    ecrire(tb, r, [None, "Période", f"={S}!E{tot}", f"={S}!G{tot}", f"={S}!F{tot}", f"={S}!I{tot}",
                   f"=F{fin_tb}-F10", f_h(r, tot),
                   f'=IFERROR(COUNTIFS({eJ},"Oui")/COUNTA({eA}),0)',
                   f'=IFERROR(COUNTIFS({dJ},"Non",{dI},">="&{P}!$C$5)/C{r},0)',
                   f"=C{r}/{S}!C{tot}", f"={S}!N{tot}", f"={S}!M{tot}"],
           {5: F_EQ, 6: F_EQ, 7: f"+{F_EQ};-{F_EQ};{F_EQ}", 9: "0%", 10: "0%", 11: "0.0", **f_res}, police=F_GRAS)
    tb.cell(r, 1).border = Border()
    for c in range(2, 14):
        tb.cell(r, c).fill = FOND_TOTAL
    tb.cell(r + 1, 2, "Évolution de la ligne « Période » : du premier au dernier mois.").font = F_NOTE
    if not reel:
        tb.conditional_formatting.add(f"H10:H{r}", DataBarRule(start_type="num", start_value=0, end_type="num",
                                                               end_value=1.3, color="6DA7EC"))
    for zone in () if reel else (f"F10:F{r} M10:M{r}",):
        tb.conditional_formatting.add(zone, FormulaRule(formula=[f"$F10>{P}!$C$3"], fill=rouge))
        tb.conditional_formatting.add(zone, FormulaRule(formula=[f"$F10<={P}!$C$3"], fill=vert))

    # Graphiques.
    g = r + 3
    graphique(wb["Synthèse"], fin_syn, g, dest=tb, colonne="B")
    graphique_tickets(wb["Synthèse"], fin_syn, tb, f"B{g + 23}")

    # Par équipe.
    r = g + 47
    tb.cell(r, 2, "Par équipe, sur toute la période").font = Font(name=POLICE, size=12, bold=True)
    r += 1
    noms = ["Équipe / technicien", "", "Jours travaillés", "Tickets faits", "Bonifications", "Tickets comptés",
            "Tickets par jour", "Journées finies après 16 h", "Part des journées après 16 h",
            "Tickets de 2 h 30 et plus", "Part des bonifications de la période"]
    for i, nom in enumerate(noms):
        c = tb.cell(r, 2 + i, nom)
        c.font, c.fill, c.border = F_GRAS, FOND_ENTETE, BORD
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    tb.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    tb.row_dimensions[r].height = 42
    debut_eq = r + 1
    compte = collections.Counter(t["equipe"] for _, ts, _ in mois for t in ts if not t["exclu"])
    eB, dC, dM = f"{EJ}!$B$4:$B${fin_ej}", f"{D}!$C$2:$C${fin_det}", f"{D}!$M$2:$M${fin_det}"
    for equipe, _ in compte.most_common():
        r += 1
        ecrire(tb, r, [None, equipe, None, f"=COUNTIFS({eB},B{r})",
                       f'=COUNTIFS({dC},B{r},{dJ},"Non")', f"=SUMIFS({dM},{dC},B{r})", f"=E{r}+F{r}",
                       f"=IFERROR(E{r}/D{r},0)", f'=COUNTIFS({eB},B{r},{eJ},"Oui")', f"=IFERROR(I{r}/D{r},0)",
                       f'=COUNTIFS({dC},B{r},{dJ},"Non",{dI},">="&{P}!$C$5)',
                       f"=IFERROR(F{r}/SUM($F${debut_eq}:$F${debut_eq + len(compte) - 1}),0)"],
               {8: "0.0", 10: "0%", 12: "0%"})
        tb.cell(r, 1).border = Border()
        tb.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    tb.conditional_formatting.add(f"E{debut_eq}:E{r}", DataBarRule(start_type="num", start_value=0,
                                                                     end_type="max", color="6DA7EC"))
    tb.conditional_formatting.add(f"F{debut_eq}:F{r}", DataBarRule(start_type="num", start_value=0,
                                                                     end_type="max", color="1C5CAB"))
    tb.page_setup.orientation = "landscape"
    tb.page_setup.fitToWidth, tb.page_setup.fitToHeight = 1, 0
    tb.sheet_properties.pageSetUpPr.fitToPage = True
    wb.active = 0


def lire_reprise(chemin):
    """Saisies d'un classeur déjà rempli : temps réels de « À vérifier » et
    bonifications tapées à la main dans « Détail tickets », par n° de ticket."""
    wb = load_workbook(chemin)
    temps, bonifs = {}, {}
    for ligne in wb["À vérifier"].iter_rows(min_row=4, values_only=True):
        if ligne[1] and isinstance(ligne[7], (int, float)):
            temps[str(ligne[1]).strip()] = ligne[7]
    for ligne in wb["Détail tickets"].iter_rows(min_row=2, values_only=True):
        if ligne[1] and isinstance(ligne[12], (int, float)):
            bonifs[str(ligne[1]).strip()] = ligne[12]
    return temps, bonifs


def construire(mois, sortie, reprise=({}, {})):
    tous = [t for _, ts, _ in mois for t in ts]
    annees = sorted({t["debut"].year for t in tous} | {t["fin"].year for t in tous})

    wb = Workbook()
    syn = wb.active
    syn.title = "Synthèse"
    par = wb.create_sheet("Paramètres")
    ej = wb.create_sheet("Par équipe et jour")
    av = wb.create_sheet("À vérifier")
    det = wb.create_sheet("Détail tickets")

    # Paramètres
    par["A1"] = "Paramètres du calcul"
    par["A1"].font = F_TITRE
    par["A2"] = "Cases jaunes à texte bleu : modifiables, tout le classeur se recalcule."
    par["A2"].font = F_NOTE
    for r, (lib, val, note) in enumerate([
        ("Forfait payé (équipes par mois)", "Au réel" if FORFAIT is None else FORFAIT,
         "Pas de forfait : payé au nombre d'équipes réel" if FORFAIT is None else "Contrat"),
        ("Tickets par équipe et par jour", TICKETS_PAR_JOUR, "Journée type 8 h - 16 h = 4 tickets"),
        ("Majoration à partir de (heures)", SEUIL_H, "Tolérance : moins de 2 h 30 = 1 ticket"),
        ("Puis +1 ticket toutes les (heures)", TRANCHE_H, "2 h 30 = 2, 4 h 30 = 3, 6 h 30 = 4…"),
        ("Fin de journée normale", FIN_JOURNEE, "Journée finie après : les longs tickets sont majorés"),
    ], start=3):
        par.cell(r, 2, lib).font = F_NORMAL
        c = par.cell(r, 3, val)
        c.font, c.fill, c.border = F_SAISIE, FOND_SAISIE, BORD
        if isinstance(val, dt.time):
            c.number_format = F_HEURE
        par.cell(r, 4, note).font = F_NOTE
    par["B8"] = "Préfixe des tickets IRIS (exclus, payés à part)"
    par["B8"].font = F_NORMAL
    par["C8"] = PREFIXE_IRIS
    par["C8"].font = F_NORMAL
    par["D8"] = "Appliqué par le programme au moment de l'import"
    par["D8"].font = F_NOTE
    if FORFAIT is None:
        par["B9"] = "Prix d'une équipe par mois (€)"
        par["B9"].font = F_NORMAL
        c = par["C9"]
        c.value, c.number_format = PRIX, F_EUR
        c.font, c.fill, c.border = F_SAISIE, FOND_SAISIE, BORD
        par["D9"] = "Montant facturé = équipes du mois × ce prix"
        par["D9"].font = F_NOTE
    par["B10"] = "Jours fériés (retirés des jours ouvrés)"
    par["B10"].font = F_GRAS
    feries = [f for a in annees for f in jours_feries(a)]
    for r, (jour, nom) in enumerate(feries, start=11):
        par.cell(r, 2, jour).number_format = F_DATE
        par.cell(r, 2).font = F_NORMAL
        par.cell(r, 3, nom).font = F_NORMAL
    plage_feries = f"{P}!$B$11:$B${10 + len(feries)}"
    for col, larg in zip("ABCD", (2, 44, 16, 46)):
        par.column_dimensions[col].width = larg

    # Détail tickets : une ligne par ticket. Les bonifications sont reliées
    # plus bas, une fois les lignes des autres onglets connues.
    entetes(det, 1, ["Mois", "N° ticket", "Équipe / technicien", "Site", "Type d'intervention",
                     "Début", "Soldée", "Jour", "Durée affichée (h)", "Exclu (IRIS / devis)",
                     "Plusieurs jours / nuit", "Ticket compté", "Bonification", "Total tickets",
                     "Pourquoi cette bonification (tickets de 2 h 30 et plus)"],
            [14, 22, 18, 18, 30, 16, 16, 11, 10, 9, 10, 9, 12, 9, 70])
    ligne_det = {}
    attendu = collections.defaultdict(int)  # bonification calculée, pour la reprise
    r = 2
    for cle, ts, _ in mois:
        for t in sorted(ts, key=lambda t: (t["equipe"], t["debut"])):
            ligne_det[id(t)] = r
            ecrire(det, r, [libelle_mois(cle), t["num"], t["equipe"], t["site"], t["type"],
                            t["debut"], t["fin"], t["debut"].date(), f"=ROUND((G{r}-F{r})*24,4)",
                            "IRIS" if t["iris"] else "Devis" if t["devis"] else "Non",
                            "Oui" if t["multi"] else "Nuit" if t["nuit"] else "Non",
                            f'=IF(J{r}="Non",1,0)', 0, f"=L{r}+M{r}"],
                   {6: F_DATEH, 7: F_DATEH, 8: F_DATE, 9: F_DEC})
            r += 1
    fin_det = r - 1

    # Par équipe et jour : tickets d'une seule journée, hors IRIS.
    ej["A1"] = "Temps occupé par équipe et par jour"
    ej["A1"].font = F_TITRE
    ej["A2"] = ("Temps occupé = de Début à Soldée des tickets du jour, chevauchements comptés une fois "
                "(calculé à l'import). Un ticket soldé le lendemain avant 6 h compte comme travail de nuit sur son jour "
                "de début. Hors IRIS et hors tickets sur plusieurs jours (voir « À vérifier »). "
                "Si le dernier ticket est soldé après la fin de journée normale (16 h), les tickets de 2 h 30 et plus "
                "sont majorés même si l'équipe a fait d'autres tickets (colonne « Temps des tickets de 2 h 30 et plus », "
                "calculée à l'import). La bonification du jour est portée par le ticket le plus long dans « Détail tickets ».")
    ej["A2"].font = F_NOTE
    entetes(ej, 3, ["Mois", "Équipe / technicien", "Jour", "", "Nb tickets faits", "Premier début",
                    "Dernière fin", "Temps occupé (h)", "Tickets selon le temps", "Fini après 16 h",
                    "Tickets de 2 h 30 et plus", "Temps des tickets de 2 h 30 et plus (h)",
                    "Tickets si les longs sont majorés", "Tickets comptés", "Bonification"],
            [14, 22, 11, 6, 9, 9, 9, 10, 10, 9, 10, 12, 11, 10, 12])
    r = 4
    for cle, ts, _ in mois:
        groupes = collections.defaultdict(list)
        for t in ts:
            if not t["exclu"] and not t["multi"]:
                groupes[(t["equipe"], t["debut"].date())].append(t)
        for (equipe, jour), g in sorted(groupes.items()):
            ecrire(ej, r, [libelle_mois(cle), equipe, jour, JOURS[jour.weekday()],
                           f'=COUNTIFS({D}!$A$2:$A${fin_det},A{r},{D}!$C$2:$C${fin_det},B{r},'
                           f'{D}!$H$2:$H${fin_det},C{r},{D}!$J$2:$J${fin_det},"Non",'
                           f'{D}!$K$2:$K${fin_det},"<>Oui")',
                           min(t["debut"] for t in g).time(), max(t["fin"] for t in g),
                           temps_occupe([(t["debut"], t["fin"]) for t in g]),
                           "=" + formule_tickets(f"H{r}"),
                           f'=IF(G{r}>C{r}+{P}!$C$7,"Oui","Non")',
                           f'=COUNTIFS({D}!$A$2:$A${fin_det},A{r},{D}!$C$2:$C${fin_det},B{r},'
                           f'{D}!$H$2:$H${fin_det},C{r},{D}!$J$2:$J${fin_det},"Non",'
                           f'{D}!$K$2:$K${fin_det},"<>Oui",{D}!$I$2:$I${fin_det},">="&{P}!$C$5)',
                           temps_occupe([(t["debut"], t["fin"]) for t in longs(g)]),
                           f"=IF(K{r}=0,E{r},E{r}-K{r}+{formule_tickets(f'L{r}')})",
                           f'=IF(J{r}="Oui",MAX(E{r},I{r},M{r}),MAX(E{r},I{r}))', f"=N{r}-E{r}"],
                   {3: F_DATE, 6: F_HEURE, 7: F_HEURE, 8: F_DEC, 12: F_DEC})
            porteur = max(g, key=lambda t: (t["fin"] - t["debut"], -ligne_det[id(t)]))
            c = det.cell(ligne_det[id(porteur)], 13, f"={EJ}!O{r}")
            c.font = F_LIEN
            attendu[id(porteur)] = tickets_journee(g) - len(g)
            journee = (f'"journée de "&ROUND({EJ}!H{r},1)&" h finie à "&TEXT({EJ}!G{r},"hh:mm")&", "'
                       f'&{EJ}!E{r}&" ticket(s) faits, "&{EJ}!N{r}&" compté(s)"')
            for t in g:
                d = ligne_det[id(t)]
                sans = f'IF(I{d}<{P}!$C$5,"","Pas de bonif : "&{journee})'
                if t is porteur:
                    expl = f'=IF(M{d}>0,"Bonif +"&M{d}&" : "&{journee},{sans})'
                else:
                    expl = (f'=IF(I{d}<{P}!$C$5,"",IF({EJ}!O{r}>0,'
                            f'"Bonif de la journée comptée sur le ticket {porteur["num"]}",{sans}))')
                det.cell(d, 15, expl).font = F_NORMAL
                det.cell(d, 15).border = BORD
            r += 1

    # À vérifier : tickets ouverts sur plusieurs jours, hors IRIS.
    av["A1"] = "Tickets ouverts sur plusieurs jours"
    av["A1"].font = F_TITRE
    av["A2"] = ("La durée affichée n'est pas du travail réel (ticket soldé plus tard). Par défaut : "
                "1 ticket sans majoration. Si vous connaissez le temps réellement travaillé, tapez-le "
                "en heures dans la colonne jaune (ex. 8 pour une intervention de 16 h à 0 h, 7.5 pour 7 h 30). "
                "Plusieurs tickets faits ensemble : mettre le temps sur un seul d'entre eux. "
                "Une ligne reste violette tant que son temps n'est pas saisi.")
    av["A2"].font = F_NOTE
    entetes(av, 3, ["Mois", "N° ticket", "Équipe / technicien", "Site", "Début", "Soldée",
                    "Durée affichée (h)", "Temps réellement travaillé (h)", "Tickets comptés",
                    "Bonification"],
            [14, 22, 18, 18, 16, 16, 10, 14, 10, 12])
    r = 4
    for cle, ts, _ in mois:
        for t in sorted((t for t in ts if t["multi"] and not t["exclu"]), key=lambda t: (t["equipe"], t["debut"])):
            ecrire(av, r, [libelle_mois(cle), t["num"], t["equipe"], t["site"], t["debut"], t["fin"],
                           f"=ROUND((F{r}-E{r})*24,4)", reprise[0].get(t["num"]),
                           f'=IF(H{r}="",1,{formule_tickets(f"H{r}")})', f"=I{r}-1"],
                   {5: F_DATEH, 6: F_DATEH, 7: F_DEC, 8: F_DEC})
            av.cell(r, 8).fill, av.cell(r, 8).font = FOND_SAISIE, F_SAISIE
            d = ligne_det[id(t)]
            h = reprise[0].get(t["num"])
            attendu[id(t)] = tickets(h) - 1 if h is not None else 0
            c = det.cell(d, 13, f"={AV}!J{r}")
            c.font = F_LIEN
            det.cell(d, 15, f'=IF(M{d}>0,"Bonif +"&M{d}&" : temps réel saisi dans « À vérifier »",'
                            f'"Sur plusieurs jours : 1 ticket, temps réel à saisir dans « À vérifier »")')
            det.cell(d, 15).font, det.cell(d, 15).border = F_NORMAL, BORD
            r += 1
    if r == 4:
        av["A4"] = "Aucun ticket sur plusieurs jours."
        av["A4"].font = F_NOTE

    # Synthèse
    syn["A1"] = (f"{NOM + ' : ' if NOM else ''}payé au réel, combien d'équipes avec la majoration ?" if FORFAIT is None
                 else f"{titre_forfait()[0].upper()}{titre_forfait()[1:]} : est-ce qu'on dépasse avec la majoration ?")
    syn["A1"].font = F_TITRE
    syn["A2"] = ("Équipes = tickets ÷ tickets par équipe et par jour ÷ jours ouvrés (lundi à vendredi hors "
                 "fériés). Les tickets IRIS (MAL-) sont exclus"
                 + (", ainsi que les tickets Cerfa payés sur devis" if CERFA_DEVIS else "")
                 + ". Majoration : voir les règles plus bas."
                 + (" Montant = équipes × prix d'une équipe par mois (onglet Paramètres). Ligne Total : "
                    "équipes en moyenne, montants additionnés." if FORFAIT is None else ""))
    syn["A2"].font = F_NOTE
    noms = ["Mois", "Premier jour", "Jours ouvrés",
            "Tickets exclus (IRIS, devis)" if CERFA_DEVIS else "Tickets IRIS (exclus)", "Tickets faits",
            "Équipes sans majoration", "Bonifications", "Tickets avec majoration",
            "Équipes avec majoration", "Forfait (équipes)", "Tickets couverts par le forfait",
            "Écart en tickets (+ = au-dessus)", "Résultat", "Tickets sur plusieurs jours à vérifier"]
    if FORFAIT is None:
        noms[9:13] = ["Apport de la majoration (équipes)", "Montant sans majoration (€)",
                      "Montant avec majoration (€)", "Gain de la majoration (€)"]
    entetes(syn, 4, noms, [16, 12, 9, 10, 9, 11, 12, 11, 11, 9, 11, 12, 22, 12])
    syn.freeze_panes = None
    dA, dJ, dK, dM = (f"{D}!${c}$2:${c}${fin_det}" for c in "AJKM")
    r = 5
    for cle, _, _ in mois:
        ecrire(syn, r, [
            libelle_mois(cle), dt.date(cle[0], cle[1], 1),
            f"=NETWORKDAYS(B{r},EOMONTH(B{r},0),{plage_feries})",
            f'=COUNTIFS({dA},A{r},{dJ},"<>Non")',
            f'=COUNTIFS({dA},A{r},{dJ},"Non")',
            f"=E{r}/{P}!$C$4/C{r}",
            f"=SUMIFS({dM},{dA},A{r})",
            f"=E{r}+G{r}",
            f"=H{r}/{P}!$C$4/C{r}",
            *([f"=I{r}-F{r}", f"=F{r}*{P}!$C$9", f"=I{r}*{P}!$C$9", f"=L{r}-K{r}"] if FORFAIT is None else
              [f"={P}!$C$3", f"=J{r}*{P}!$C$4*C{r}", f"=H{r}-K{r}",
               f'=IF(I{r}>J{r},"AU-DESSUS du forfait","En dessous du forfait")']),
            f'=COUNTIFS({dA},A{r},{dK},"Oui",{dJ},"Non")',
        ], {2: F_DATE, 6: F_EQ, 9: F_EQ, 10: F_EQ,
            **({11: F_EUR, 12: F_EUR, 13: F_EUR} if FORFAIT is None else {11: "0.0", 12: "+0.0;-0.0;0"})})
        r += 1
    fin_syn = r - 1
    ecrire(syn, r, ["Total / moyenne", None, f"=SUM(C5:C{fin_syn})", f"=SUM(D5:D{fin_syn})",
                    f"=SUM(E5:E{fin_syn})", f"=E{r}/{P}!$C$4/C{r}", f"=SUM(G5:G{fin_syn})",
                    f"=SUM(H5:H{fin_syn})", f"=H{r}/{P}!$C$4/C{r}",
                    *([f"=I{r}-F{r}", f"=SUM(K5:K{fin_syn})", f"=SUM(L5:L{fin_syn})", f"=SUM(M5:M{fin_syn})"]
                      if FORFAIT is None else
                      [f"={P}!$C$3", f"=SUM(K5:K{fin_syn})", f"=H{r}-K{r}",
                       f'=IF(I{r}>J{r},"AU-DESSUS du forfait","En dessous du forfait")']),
                    f"=SUM(N5:N{fin_syn})"],
           {6: F_EQ, 9: F_EQ, 10: F_EQ,
            **({11: F_EUR, 12: F_EUR, 13: F_EUR} if FORFAIT is None else {11: "0.0", 12: "+0.0;-0.0;0"})},
           police=F_GRAS)
    for c in range(1, 15):
        syn.cell(r, c).fill = FOND_TOTAL
    rouge = PatternFill("solid", fgColor="F8D7DA")
    vert = PatternFill("solid", fgColor="D4EDDA")
    if FORFAIT is not None:
        zone = f"I5:I{r} M5:M{r}"
        syn.conditional_formatting.add(zone, FormulaRule(formula=[f"$I5>$J5"], fill=rouge))
        syn.conditional_formatting.add(zone, FormulaRule(formula=[f"$I5<=$J5"], fill=vert))

    regles = [
        "Règles du calcul",
        "• Pour chaque équipe et chaque jour : temps occupé par ses tickets (chevauchements comptés une fois).",
        "• Moins de 2 h 30 = 1 ticket ; 2 h 30 = 2 ; 4 h 30 = 3 ; 6 h 30 = 4 ; puis +1 toutes les 2 h.",
        "• Jamais moins que le nombre de tickets faits : deux tickets de 2 h 30 le même jour = 3 tickets (un seul majoré).",
        "• Journée finie après 16 h : les tickets de 2 h 30 et plus sont majorés même si l'équipe a fait d'autres tickets.",
        "• Ticket soldé le lendemain avant 6 h : travail de nuit, compté en entier sur son jour de début.",
        "• Autre ticket ouvert sur plusieurs jours : 1 ticket par défaut, temps réel à saisir dans « À vérifier ».",
        "• Week-ends et jours fériés : les tickets comptent et sont majorés, mais ces jours ne comptent pas dans les jours ouvrés.",
        "• Tickets MAL- (IRIS) : payés à part, exclus.",
        *(["• Tickets avec Cerfa / Plainte / Constat = Oui : payés sur devis, exclus."] if CERFA_DEVIS else []),
        "• Un nom porté par plusieurs techniciens (ex. [RES] SADOK) est séparé par prénom. Sans prénom dans "
        "l'extraction, SADOK est séparé en A et B d'après l'enchaînement des horaires.",
        "• Dans « Détail tickets » et « Par équipe et jour » : lignes violettes = à vérifier, lignes orange = 2 h 30 ou plus, case verte = bonification.",
    ]
    for i, txt in enumerate(regles):
        c = syn.cell(r + 2 + i, 1, txt)
        c.font = F_GRAS if i == 0 else F_NORMAL

    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    orange = PatternFill("solid", fgColor="FCE4D6")
    violet = PatternFill("solid", fgColor="E4DFEC")
    det.conditional_formatting.add(f"M2:M{fin_det}", FormulaRule(formula=["$M2>0"], fill=vert))
    det.conditional_formatting.add(f"A2:O{fin_det}", FormulaRule(formula=['$K2="Oui"'], fill=violet))
    det.conditional_formatting.add(f"A2:O{fin_det}", FormulaRule(formula=[f"$I2>={P}!$C$5"], fill=orange))
    # Bonifications tapées à la main dans un ancien classeur : gardées telles quelles.
    for cle, ts, _ in mois:
        for t in ts:
            d = ligne_det[id(t)]
            avant = reprise[1].get(t["num"])
            if avant is None or avant == attendu[id(t)]:
                continue
            det.cell(d, 13, avant).font = F_SAISIE
            det.cell(d, 13).comment = Comment("Bonification saisie à la main (reprise de l'ancien classeur)", "Forfait")
            det.cell(d, 15, f'="Bonif +"&M{d}&" saisie à la main"')
    fin_ej = max(ej.max_row, 4)
    ej.conditional_formatting.add(f"O4:O{fin_ej}", FormulaRule(formula=["$O4>0"], fill=vert))
    ej.conditional_formatting.add(f"A4:O{fin_ej}", FormulaRule(formula=[f"$H4>={P}!$C$5"], fill=orange))
    fin_av = max(av.max_row, 4)
    for zone in (f"A4:G{fin_av}", f"I4:J{fin_av}"):
        av.conditional_formatting.add(zone, FormulaRule(formula=['AND($B4<>"",$H4="")'], fill=violet))
    graphique(syn, fin_syn, r + 2 + len(regles) + 1)
    tableau_de_bord(wb, mois, fin_syn, fin_det, fin_ej)
    det.auto_filter.ref = f"A1:O{fin_det}"
    ej.auto_filter.ref = f"A3:O{ej.max_row}"
    av.auto_filter.ref = f"A3:J{av.max_row}"
    wb.save(sortie)


def main():
    global FORFAIT, NOM, F_EQ, PRIX, CERFA_DEVIS
    ici = os.path.dirname(os.path.abspath(__file__))
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("exports", nargs="+", help="Extractions Praxedo (.csv ou .xlsx), une par mois")
    ap.add_argument("-o", "--sortie", help="Classeur à écrire (par défaut Forfait_equipes[_NOM].xlsx à côté du programme)")
    ap.add_argument("--forfait", type=float, default=FORFAIT, help="Forfait payé, en équipes par mois (3.8 par défaut)")
    ap.add_argument("--nom", default="", help="Nom du contrat, affiché dans les titres (ex. Glaucos)")
    ap.add_argument("--reel", action="store_true",
                    help="Pas de forfait, payé au réel : équipes du mois × prix d'une équipe par mois")
    ap.add_argument("--cerfa-devis", action="store_true",
                    help="Exclure les tickets Cerfa / Plainte / Constat = Oui (payés sur devis)")
    ap.add_argument("--prix", type=float, help="Payé au réel : prix d'une équipe par mois, en euros (modifiable ensuite)")
    ap.add_argument("--reprendre", metavar="ANCIEN.xlsx",
                    help="Classeur déjà rempli : reprend ses temps « À vérifier » et ses bonifications tapées à la main")
    args = ap.parse_args()
    FORFAIT, NOM, PRIX = (None if args.reel else args.forfait), args.nom.strip(), args.prix
    CERFA_DEVIS = args.cerfa_devis
    F_EQ = "0.000" if FORFAIT is not None and FORFAIT < 1 else F_DEC
    args.sortie = args.sortie or os.path.join(ici, f"Forfait_equipes{'_' + NOM if NOM else ''}.xlsx")
    reprise = lire_reprise(args.reprendre) if args.reprendre else ({}, {})

    par_mois = {}
    for chemin in args.exports:
        cle, tickets, ignores = lire_export(chemin)
        deja = par_mois.setdefault(cle, ([], []))
        vus = {t["num"] for t in deja[0]}
        deja[0].extend(t for t in tickets if t["num"] not in vus)
        deja[1].extend(ignores)
        print(f"{os.path.basename(chemin)} : {libelle_mois(cle)}, {len(tickets)} tickets"
              + (f", {len(ignores)} ignorés (annulés ou sans dates)" if ignores else ""))
    # Un mois sans extraction entre le premier et le dernier compte 0 ticket.
    a, m = min(par_mois)
    while (a, m) < max(par_mois):
        if (a, m) not in par_mois:
            par_mois[(a, m)] = ([], [])
            print(f"{libelle_mois((a, m))} : aucune extraction, compté à 0 ticket")
        a, m = (a + 1, 1) if m == 12 else (a, m + 1)
    mois = [(cle, ts, ig) for cle, (ts, ig) in sorted(par_mois.items())]
    for _, ts, _ in mois:
        nommer_par_prenom(ts)
        separer_equipes(ts)
    construire(mois, args.sortie, reprise)
    print(f"Classeur écrit : {args.sortie}")


if __name__ == "__main__":
    main()
