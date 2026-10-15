"""Condensé de plusieurs classeurs « Forfait équipes » (un par projet) face
à un forfait global.

    python3 forfait-equipes/synthese_projets.py --forfait 4 \
        GNY=Forfait_equipes_GNY.xlsx Glaucos=Forfait_equipes_Glaucos.xlsx Resoptic=Forfait_equipes_Resoptic.xlsx

Chaque classeur doit avoir été enregistré par Excel (ou recalculé) : on lit
les valeurs de son onglet « Synthèse » (tickets faits, bonifications, jours
ouvrés, tickets à vérifier) et son forfait dans « Paramètres ». Le condensé
recalcule tout le reste par formules.
"""

import argparse
import datetime as dt
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.chart import BarChart, LineChart, Reference  # noqa: E402
from openpyxl.chart.label import DataLabelList  # noqa: E402
from openpyxl.chart.shapes import GraphicalProperties  # noqa: E402
from openpyxl.drawing.line import LineProperties  # noqa: E402
from openpyxl.formatting.rule import DataBarRule, FormulaRule  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

import calcul_forfait as cf  # noqa: E402

# Couleurs des projets, dans l'ordre (palette catégorielle validée).
COULEURS = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4"]
P, DO = "'Paramètres'", "'Données'"


def lire_projet(chemin):
    wb = load_workbook(chemin, data_only=True)
    forfait = wb["Paramètres"]["C3"].value
    mois = []
    for ligne in wb["Synthèse"].iter_rows(min_row=5, values_only=True):
        if not ligne[0] or ligne[0] == "Total / moyenne" or ligne[2] is None:
            break
        premier = ligne[1].date() if isinstance(ligne[1], dt.datetime) else ligne[1]
        mois.append({"mois": ligne[0], "premier": premier, "jours": ligne[2], "faits": ligne[4],
                     "bonifs": ligne[6], "verifier": ligne[13]})
    if not mois or any(m["faits"] is None for m in mois):
        sys.exit(f"{chemin} : pas de valeurs calculées dans « Synthèse ». Ouvrir et enregistrer le fichier dans Excel.")
    return forfait, mois


def tuile(ws, col, titre, valeur, sous, fmt=None):
    fond = PatternFill("solid", fgColor="F3F6FA")
    c1, c2 = get_column_letter(col), get_column_letter(col + 1)
    for r, contenu, police in ((4, titre, Font(name=cf.POLICE, size=9, color="52514E")),
                               (5, valeur, Font(name=cf.POLICE, size=20 if fmt else 15, bold=True)),
                               (6, sous, Font(name=cf.POLICE, size=9, color="52514E"))):
        ws.merge_cells(f"{c1}{r}:{c2}{r}")
        c = ws[f"{c1}{r}"]
        c.value, c.font = contenu, police
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if fmt and r == 5:
            c.number_format = fmt
        for col_ in (c1, c2):
            ws[f"{col_}{r}"].fill = fond


def entetes(ws, ligne, col, noms):
    for i, nom in enumerate(noms):
        c = ws.cell(ligne, col + i, nom)
        c.font, c.fill, c.border = cf.F_GRAS, cf.FOND_ENTETE, cf.BORD
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[ligne].height = 42


def construire(projets, forfait_global, sortie):
    wb = Workbook()
    tb = wb.active
    tb.title = "Tableau de bord"
    par = wb.create_sheet("Paramètres")
    don = wb.create_sheet("Données")
    noms = [nom for nom, _, _ in projets]
    tous_mois = []
    for _, _, mois in projets:
        for m in mois:
            if m["mois"] not in [x["mois"] for x in tous_mois]:
                tous_mois.append(m)
    tous_mois.sort(key=lambda m: m["premier"])

    # Paramètres
    par["A1"] = "Paramètres du condensé"
    par["A1"].font = cf.F_TITRE
    par["A2"] = "Cases jaunes à texte bleu : modifiables, tout le classeur se recalcule."
    par["A2"].font = cf.F_NOTE
    somme = sum(f for _, f, _ in projets)
    for r, (lib, val, note) in enumerate([
        ("Forfait global payé (équipes par mois)", forfait_global,
         f"Somme des forfaits des projets : {' + '.join(f'{f:g}' for _, f, _ in projets)} = {somme:g}".replace(".", ",")),
        ("Tickets par équipe et par jour", cf.TICKETS_PAR_JOUR, "Journée type 8 h - 16 h = 4 tickets"),
    ], start=3):
        par.cell(r, 2, lib).font = cf.F_NORMAL
        c = par.cell(r, 3, val)
        c.font, c.fill, c.border = cf.F_SAISIE, cf.FOND_SAISIE, cf.BORD
        par.cell(r, 4, note).font = cf.F_NOTE
    par["B6"] = "Forfait de chaque projet"
    par["B6"].font = cf.F_GRAS
    for i, (nom, forfait, _) in enumerate(projets):
        par.cell(7 + i, 2, nom).font = cf.F_NORMAL
        c = par.cell(7 + i, 3, forfait)
        c.font, c.fill, c.border = cf.F_SAISIE, cf.FOND_SAISIE, cf.BORD
    ligne_forfait = {nom: 7 + i for i, (nom, _, _) in enumerate(projets)}
    for col, larg in zip("ABCD", (2, 40, 12, 50)):
        par.column_dimensions[col].width = larg

    # Données : une ligne par projet et par mois, reprises des classeurs projet.
    don["A1"] = "Données reprises des classeurs de chaque projet (onglet « Synthèse »)"
    don["A1"].font = cf.F_TITRE
    entetes(don, 3, 1, ["Projet", "Mois", "Premier jour", "Jours ouvrés", "Tickets faits", "Bonifications",
                        "Tickets avec majoration", "Équipes sans majoration", "Équipes avec majoration",
                        "Forfait du projet", "Au-dessus du forfait du projet", "Tickets à vérifier"])
    for col, larg in zip("ABCDEFGHIJKL", (12, 15, 12, 9, 9, 12, 11, 11, 11, 10, 11, 10)):
        don.column_dimensions[col].width = larg
    don.freeze_panes = "A4"
    r = 4
    for nom, _, mois in projets:
        par_mois = {m["mois"]: m for m in mois}
        for m in tous_mois:
            x = par_mois.get(m["mois"], {"faits": 0, "bonifs": 0, "verifier": 0})
            cf.ecrire(don, r, [nom, m["mois"], m["premier"], m["jours"], x["faits"], x["bonifs"],
                               f"=E{r}+F{r}", f"=E{r}/{P}!$C$4/D{r}", f"=G{r}/{P}!$C$4/D{r}",
                               f"={P}!$C${ligne_forfait[nom]}", f"=IF(I{r}>J{r},1,0)", x["verifier"]],
                      {3: cf.F_DATE, 8: "0.000", 9: "0.000", 10: "0.00"})
            r += 1
    fin_don = r - 1
    dA, dB, dD, dE, dF, dG, dH, dI, dK, dL = (f"{DO}!${c}$4:${c}${fin_don}" for c in "ABDEFGHIKL")

    # Tableau de bord
    tb.sheet_view.showGridLines = False
    tb.column_dimensions["A"].width = 2
    for i in range(2, 15):
        tb.column_dimensions[get_column_letter(i)].width = 12.5
    tb.column_dimensions["B"].width = 16
    tb["B1"] = f"Synthèse {len(projets)} projets : forfait global {forfait_global:g} équipes".replace(".", ",")
    tb["B1"].font = Font(name=cf.POLICE, size=16, bold=True)
    tb["B2"] = (f'="{", ".join(noms)} · période : "&B19&" à "&B{18 + len(tous_mois)}'
                f'&" · forfait global "&FIXED({P}!$C$3,2)&" équipes par mois"')
    tb["B2"].font = cf.F_NOTE

    # Mois par mois (lignes 18 et suivantes), construit avant les tuiles qui s'y réfèrent.
    m0 = 19
    m1 = m0 + len(tous_mois) - 1
    tot = m1 + 1
    tb.cell(17, 2, "Mois par mois, tous projets (équipes avec majoration)").font = Font(name=cf.POLICE, size=12, bold=True)
    cols = ["Mois", "Jours ouvrés"] + noms + ["Total avec majoration", "Total sans majoration",
                                             "Forfait global", "Écart", "Part du forfait utilisée", "Résultat"]
    entetes(tb, 18, 2, cols)
    k = len(noms)
    cT, cS, cF, cE, cP, cR = (get_column_letter(4 + k + i) for i in range(6))
    for i, m in enumerate(tous_mois):
        r = m0 + i
        vals = [None, m["mois"], f"=INDEX({dD},MATCH(B{r},{dB},0))"]
        vals += [f'=SUMIFS({dI},{dA},"{nom}",{dB},$B{r})' for nom in noms]
        vals += [f"=SUM(D{r}:{get_column_letter(3 + k)}{r})", f"=SUMIFS({dH},{dB},$B{r})", f"={P}!$C$3",
                 f"={cT}{r}-{cF}{r}", f"={cT}{r}/{cF}{r}",
                 f'=IF({cT}{r}>{cF}{r},"AU-DESSUS du forfait","En dessous du forfait")']
        fm = {4 + j: "0.000" for j in range(k)}
        fm.update({4 + k: "0.00", 5 + k: "0.00", 6 + k: "0.00", 7 + k: "+0.00;-0.00;0.00", 8 + k: "0%"})
        cf.ecrire(tb, r, vals, fm)
        tb.cell(r, 1).border = cf.Border()
    vals = [None, "Période", f"=SUM(C{m0}:C{m1})"]
    vals += [f'=SUMIFS({dG},{dA},"{nom}")/{P}!$C$4/C{tot}' for nom in noms]
    vals += [f"=SUM(D{tot}:{get_column_letter(3 + k)}{tot})", f"=SUM({dE})/{P}!$C$4/C{tot}", f"={P}!$C$3",
             f"={cT}{tot}-{cF}{tot}", f"={cT}{tot}/{cF}{tot}",
             f'=IF({cT}{tot}>{cF}{tot},"AU-DESSUS du forfait","En dessous du forfait")']
    fm = {4 + j: "0.000" for j in range(k)}
    fm.update({4 + k: "0.00", 5 + k: "0.00", 6 + k: "0.00", 7 + k: "+0.00;-0.00;0.00", 8 + k: "0%"})
    cf.ecrire(tb, tot, vals, fm, police=cf.F_GRAS)
    tb.cell(tot, 1).border = cf.Border()
    for c in range(2, 10 + k):
        tb.cell(tot, c).fill = cf.FOND_TOTAL
    tb.column_dimensions[cR].width = 22
    rouge, vert = PatternFill("solid", fgColor="F8D7DA"), PatternFill("solid", fgColor="D4EDDA")
    zone = f"{cT}{m0}:{cT}{tot} {cR}{m0}:{cR}{tot}"
    tb.conditional_formatting.add(zone, FormulaRule(formula=[f"${cT}{m0}>${cF}{m0}"], fill=rouge))
    tb.conditional_formatting.add(zone, FormulaRule(formula=[f"${cT}{m0}<=${cF}{m0}"], fill=vert))
    tb.conditional_formatting.add(f"{cP}{m0}:{cP}{tot}", DataBarRule(start_type="num", start_value=0, end_type="num",
                                                                     end_value=1.3, color="6DA7EC"))

    # Par projet (lignes 9 à 9 + k).
    tb.cell(8, 2, "Par projet, sur toute la période").font = Font(name=cf.POLICE, size=12, bold=True)
    entetes(tb, 9, 2, ["Projet", "Forfait payé", "Équipes sans majoration", "Équipes avec majoration",
                       "Écart avec le forfait", "Part de son forfait utilisée", "Mois au-dessus de son forfait",
                       "Tickets faits", "Bonifications", "Tickets à vérifier", "Part de l'activité totale"])
    p0 = 10
    for i, nom in enumerate(noms):
        r = p0 + i
        cf.ecrire(tb, r, [None, nom, f"={P}!$C${ligne_forfait[nom]}",
                          f'=SUMIFS({dE},{dA},B{r})/{P}!$C$4/$C${tot}',
                          f'=SUMIFS({dG},{dA},B{r})/{P}!$C$4/$C${tot}', f"=E{r}-C{r}", f"=E{r}/C{r}",
                          f'=SUMIFS({dK},{dA},B{r})&" sur "&COUNTIFS({dA},B{r})',
                          f"=SUMIFS({dE},{dA},B{r})", f"=SUMIFS({dF},{dA},B{r})", f"=SUMIFS({dL},{dA},B{r})",
                          f"=E{r}/$E${p0 + k}"],
                  {3: "0.00", 4: "0.000", 5: "0.000", 6: "+0.000;-0.000;0.000", 7: "0%", 12: "0%"})
        tb.cell(r, 1).border = cf.Border()
    r = p0 + k
    cf.ecrire(tb, r, [None, "Total", f"={P}!$C$3", f"=SUM(D{p0}:D{r - 1})", f"=SUM(E{p0}:E{r - 1})",
                      f"=E{r}-C{r}", f"=E{r}/C{r}", f'=COUNTIF({cT}{m0}:{cT}{m1},">"&{P}!$C$3)&" sur "&{len(tous_mois)}',
                      f"=SUM(I{p0}:I{r - 1})", f"=SUM(J{p0}:J{r - 1})", f"=SUM(K{p0}:K{r - 1})", "=1"],
              {3: "0.00", 4: "0.00", 5: "0.00", 6: "+0.00;-0.00;0.00", 7: "0%", 12: "0%"}, police=cf.F_GRAS)
    tb.cell(r, 1).border = cf.Border()
    for c in range(2, 13):
        tb.cell(r, c).fill = cf.FOND_TOTAL
    tb.cell(r + 1, 2, "Total : forfait global et mois où l'ensemble des projets dépasse le forfait global.").font = cf.F_NOTE
    for zone, f in ((f"F{p0}:G{r}", f"$E{p0}>$C{p0}"),):
        tb.conditional_formatting.add(zone, FormulaRule(formula=[f], fill=rouge))
        tb.conditional_formatting.add(zone, FormulaRule(formula=[f"NOT({f})"], fill=vert))
    tb.conditional_formatting.add(f"G{p0}:G{r - 1}", DataBarRule(start_type="num", start_value=0, end_type="max",
                                                                  color="6DA7EC"))

    # Chiffres clés.
    PT = f"$B${m0}:$B${m1}"
    tuile(tb, 2, f"Équipes en moyenne, {k} projets avec majoration", f"={cT}{tot}",
          f'="sans majoration : "&FIXED({cS}{tot},2)&" · forfait "&FIXED({P}!$C$3,2)', "0.00")
    tuile(tb, 4, "Mois au-dessus du forfait global", f'=COUNTIF({cT}{m0}:{cT}{m1},">"&{P}!$C$3)&" sur {len(tous_mois)}"',
          f'="forfait global : "&FIXED({P}!$C$3,2)&" équipes"')
    tuile(tb, 6, "Mois le plus chargé", f"=INDEX({PT},MATCH(MAX({cT}{m0}:{cT}{m1}),{cT}{m0}:{cT}{m1},0))",
          f'=FIXED(MAX({cT}{m0}:{cT}{m1}),2)&" équipes, soit "&FIXED(MAX({cT}{m0}:{cT}{m1})/{P}!$C$3*100,0)&" % du forfait"')
    tuile(tb, 8, "Apport de la majoration", f'="+"&FIXED({cT}{tot}-{cS}{tot},2)&" équipe"',
          f'="+"&J{p0 + k}&" tickets, soit +"&FIXED(J{p0 + k}/I{p0 + k}*100,0)&" %"')
    tuile(tb, 10, "Projets au-dessus de leur forfait", f'=COUNTIF(F{p0}:F{p0 + k - 1},">0")&" sur {k}"',
          f'="le plus au-dessus : "&INDEX(B{p0}:B{p0 + k - 1},MATCH(MAX(G{p0}:G{p0 + k - 1}),G{p0}:G{p0 + k - 1},0))'
          f'&" ("&FIXED(MAX(G{p0}:G{p0 + k - 1})*100,0)&" %)"')
    tuile(tb, 12, "Marge sur la période (tickets)", f"=({cF}{tot}-{cT}{tot})*{P}!$C$4*C{tot}",
          f'=IF({cT}{tot}<{cF}{tot},"encore sous le forfait global","au-dessus du forfait global")', "0;-0;0")
    tb.row_dimensions[4].height, tb.row_dimensions[5].height, tb.row_dimensions[6].height = 26, 34, 30
    for zone, f in (("B5:C5", f"$B$5>{P}!$C$3"),):
        tb.conditional_formatting.add(zone, FormulaRule(formula=[f], fill=rouge))
        tb.conditional_formatting.add(zone, FormulaRule(formula=[f"NOT({f})"], fill=vert))

    # Graphique 1 : équipes par mois empilées par projet, face au forfait global.
    g = tot + 3
    barres = BarChart()
    barres.type, barres.grouping, barres.overlap, barres.gapWidth = "col", "stacked", 100, 60
    barres.title = "Équipes par mois et par projet (avec majoration), face au forfait global"
    barres.y_axis.title = "Équipes"
    barres.y_axis.scaling.min = 0
    barres.y_axis.numFmt = "0.0"
    barres.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="E6E6E3"))
    for j in range(k):
        barres.add_data(Reference(tb, min_col=4 + j, min_row=18, max_row=m1), titles_from_data=True)
        s = barres.series[-1]
        s.graphicalProperties.solidFill = COULEURS[j % len(COULEURS)]
        s.graphicalProperties.line.solidFill = "FFFFFF"
    barres.set_categories(Reference(tb, min_col=2, min_row=m0, max_row=m1))
    ligne = LineChart()
    ligne.add_data(Reference(tb, min_col=6 + k, min_row=18, max_row=m1), titles_from_data=True)
    s = ligne.series[0]
    s.graphicalProperties.line.solidFill = "52514E"
    s.graphicalProperties.line.width = 22000
    s.graphicalProperties.line.prstDash = "dash"
    s.smooth = False
    ligne.y_axis.axId = barres.y_axis.axId
    barres += ligne
    barres.x_axis.delete = barres.y_axis.delete = False
    barres.legend.position = "b"
    barres.width, barres.height = 24, 11
    tb.add_chart(barres, f"B{g}")

    # Graphique 2 : part de son forfait utilisée, par projet.
    parts = BarChart()
    parts.type, parts.gapWidth = "bar", 80
    parts.title = "Part de son forfait utilisée, par projet (100 % = forfait)"
    parts.add_data(Reference(tb, min_col=7, min_row=9, max_row=p0 + k - 1), titles_from_data=True)
    parts.set_categories(Reference(tb, min_col=2, min_row=p0, max_row=p0 + k - 1))
    s = parts.series[0]
    s.graphicalProperties.solidFill = "1C5CAB"
    s.graphicalProperties.line.noFill = True
    s.dLbls = DataLabelList()
    s.dLbls.showVal, s.dLbls.numFmt = True, "0%"
    for attr in ("showSerName", "showCatName", "showLegendKey", "showPercent"):
        setattr(s.dLbls, attr, False)
    parts.x_axis.delete = parts.y_axis.delete = False
    parts.y_axis.scaling.min = 0
    parts.y_axis.numFmt = "0%"
    parts.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="E6E6E3"))
    parts.legend = None
    parts.width, parts.height = 24, 7
    tb.add_chart(parts, f"B{g + 23}")

    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(sortie)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("projets", nargs="+", metavar="NOM=classeur.xlsx")
    ap.add_argument("--forfait", type=float, required=True, help="Forfait global payé, en équipes par mois")
    ap.add_argument("-o", "--sortie", default=os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                                           "Synthese_projets.xlsx"))
    args = ap.parse_args()
    projets = []
    for p in args.projets:
        nom, _, chemin = p.partition("=")
        forfait, mois = lire_projet(chemin)
        projets.append((nom.strip(), forfait, mois))
        print(f"{nom} : forfait {forfait:g}, {len(mois)} mois, {sum(m['faits'] for m in mois)} tickets")
    construire(projets, args.forfait, args.sortie)
    print(f"Classeur écrit : {args.sortie}")


if __name__ == "__main__":
    main()
