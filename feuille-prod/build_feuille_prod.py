# Construit Feuille_prod.xlsx : la production des techniciens comptée en journées, pas en euros.
# Chaque prestation vaut une part de journée selon le nombre qu'un technicien en fait dans une
# journée normale (onglet Règles). Les prix viennent du BPU de journee-agent/data.json.
# Usage : python3 feuille-prod/build_feuille_prod.py
import json, os
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import DataBarRule

ICI = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(ICI, '..', 'journee-agent', 'data.json')
OUT = os.path.join(ICI, 'Feuille_prod.xlsx')
data = json.load(open(DATA, encoding='utf-8'))

# Règles : (catégorie, ce qu'on compte, comment ça compte, combien par journée, remarque)
# « au prix » : la valeur d'une journée est celle d'une journée de tickets du même marché (colonne E).
REGLES = [
    ("AXIONE · City Fast",         "intervention", "au nombre", 4,
     "D'après toi : 4 par jour, pour N1, N2 et N3. Si un N3 (tirage de câble) prend plus longtemps, on peut le mettre à part."),
    ("Heure de technicien",        "heure",        "au nombre", 7,
     "AXIONE RHD57 et heures COVAGE (56.4, 63.1, 63.2). Journée de 7 h : mets 8 si vos journées font 8 h. « 2 techniciens + nacelle » : seul celui qui saisit la ligne est compté."),
    ("COVAGE · Ticket",            "ticket",       "au nombre", 2, "D'après toi : 2 par jour (sites éloignés)."),
    ("COVAGE · Autres prestations", "prestation",  "au prix",   None,
     "Raccordements, câble, boîtes, études… comptés au prix : 1 journée = ce que rapporte une journée de tickets COVAGE. Tu peux taper une autre valeur en E."),
    ("ALTITUDE · Ticket",          "ticket",       "au nombre", 4, "D'après toi : 4 par jour."),
    ("ALTITUDE · Malfaçons",       "malfaçon",     "au prix",   None,
     "Au prix d'une journée de tickets ALTITUDE : une malfaçon aérienne (200 €) ≈ 1 ticket, une souterraine (100 €) ≈ ½ ticket. À vérifier."),
    ("LOSANGE · Ticket",           "ticket",       "au nombre", 4, "D'après toi : 4 par jour."),
    ("LOSANGE · Malfaçons",        "malfaçon",     "au prix",   None, "Comme ALTITUDE, au prix d'une journée de tickets LOSANGE. À vérifier."),
    ("IELO · Ticket",              "ticket",       "au nombre", 2, "D'après toi : 2 par jour (sites éloignés)."),
    ("CELESTE · SAV",              "SAV",          "au nombre", 2,
     "D'après toi : 2 par jour, pour F0, F1 et F2. Si un F2 prend plus longtemps, on peut le mettre à part."),
    ("FREE · Ticket",              "ticket",       "au nombre", None, "À compléter : combien de tickets FREE par jour ? Le prix du ticket FREE manque aussi dans le BPU."),
    ("Ne compte pas",              "—",            "ne compte pas", None,
     "Location de nacelle, plus-values de câble, rendez-vous, majoration hotline : de l'argent en plus, pas du travail en plus."),
    ("Absence",                    "—",            "absence",   None, "Congés et absences : le jour n'est pas compté comme travaillé, il ne fait donc pas baisser le taux."),
]
# Pour les règles « au prix » : le ticket du même marché qui donne la valeur d'une journée
REF_AU_PRIX = {"COVAGE · Autres prestations": ("COVAGE · Ticket", "COVAGE-TICKET"),
               "ALTITUDE · Malfaçons": ("ALTITUDE · Ticket", "ALTITUDE-TICKET"),
               "LOSANGE · Malfaçons": ("LOSANGE · Ticket", "LOSANGE-TICKET")}
HEURES = {'MC-T1', 'MC-T2', 'MC-T3', 'MC-T4', 'MC-T5', 'COV-56.4', 'COV-63.1', 'COV-63.2'}
NE_COMPTE_PAS = {'MC-T6', 'COV-59.2', 'COV-60.3', 'COV-61.4', 'COV-12.11'} | {f'COV-55.{n}' for n in range(5, 14)}

def categorie(projet, code):
    if code in HEURES: return "Heure de technicien"
    if code in NE_COMPTE_PAS: return "Ne compte pas"
    if projet in ('CONGÉS', 'ABSENCE'): return "Absence"
    if code.endswith('-TICKET'): return f"{projet} · Ticket"
    if projet == 'AXIONE' and code.startswith('CF-'): return "AXIONE · City Fast"
    if projet in ('ALTITUDE', 'LOSANGE') and '-MALF-' in code: return f"{projet} · Malfaçons"
    if projet == 'CELESTE' and code.startswith('SAV-'): return "CELESTE · SAV"
    if projet == 'COVAGE': return "COVAGE · Autres prestations"
    raise SystemExit(f'Prestation sans catégorie : {projet} {code}')

lignes = [(p, r[0], r[1], r[3], r[4], categorie(p, r[0])) for p, rows in data['bpu'].items() for r in rows]
assert {l[5] for l in lignes} <= {r[0] for r in REGLES}
marches = [p['nom'] for p in data['projets'] if not p['absence']]

# Plages
RG1, RG2 = 5, 4 + len(REGLES)
PR1, PR2 = 5, 4 + len(lignes)
SA1, SA2 = 2, 5001                # lignes de saisie préparées
TE1, TE2 = 7, 21                  # 15 places de techniciens
R_A = f"'Règles'!$A${RG1}:$A${RG2}"; R_C = f"'Règles'!$C${RG1}:$C${RG2}"
R_D = f"'Règles'!$D${RG1}:$D${RG2}"; R_E = f"'Règles'!$E${RG1}:$E${RG2}"
P_B = f"'Prestations'!$B${PR1}:$B${PR2}"; P_C = f"'Prestations'!$C${PR1}:$C${PR2}"
P_E = f"'Prestations'!$E${PR1}:$E${PR2}"; P_F = f"'Prestations'!$F${PR1}:$F${PR2}"
P_G = f"'Prestations'!$G${PR1}:$G${PR2}"

# ---------- styles (mêmes conventions que le barème) ----------
F = lambda **k: Font(name='Arial', size=k.pop('size', 10), **k)
TITRE = F(size=14, bold=True, color='1F3864'); SOUS = F(bold=True, size=11, color='1F3864')
ENTETE = F(bold=True, color='FFFFFF'); FOND_ENTETE = PatternFill('solid', fgColor='1F3864')
SAISIE = PatternFill('solid', fgColor='FFF2CC'); BLEU = F(color='0000FF')
GRIS = PatternFill('solid', fgColor='F2F2F2'); NOTE = F(italic=True, color='595959')
fin = Side(style='thin', color='BFBFBF'); BORD = Border(left=fin, right=fin, top=fin, bottom=fin)
WRAP = Alignment(wrap_text=True, vertical='top')
EUROS, EUROS2, JOURS, DATE = '#,##0 "€"', '#,##0.00 "€"', '0.00', 'dd/mm/yyyy'

def entetes(ws, ligne, noms, hauteur=None):
    for j, h in enumerate(noms, 1):
        c = ws.cell(ligne, j, h); c.font = ENTETE; c.fill = FOND_ENTETE; c.border = BORD
        c.alignment = Alignment(wrap_text=True, vertical='center')
    if hauteur: ws.row_dimensions[ligne].height = hauteur

def largeurs(ws, cols):
    for col, w in cols.items(): ws.column_dimensions[col].width = w

wb = openpyxl.Workbook()

# ---------- Mode d'emploi ----------
ws = wb.active; ws.title = "Mode d'emploi"
ws['A1'] = 'FEUILLE DE PROD : CHAQUE TECHNICIEN JUGÉ SUR SON TRAVAIL, PAS SUR LE PRIX DU MARCHÉ'; ws['A1'].font = TITRE
texte = [
    ("LE PROBLÈME", True),
    ("Un ticket rapporte 400 € chez COVAGE et 207 € chez LOSANGE, une heure AXIONE environ 35 € (prix du BPU). Compter la production en euros, c'est juger le technicien sur le marché où tu l'envoies.", False),
    ("", False),
    ("LA SOLUTION : COMPTER EN JOURNÉES", True),
    ("Pour chaque marché, tu dis combien un technicien en fait dans une journée normale (onglet Règles) : 4 tickets LOSANGE = 1 journée, 2 tickets COVAGE = 1 journée, 7 heures AXIONE = 1 journée.", False),
    ("Taux de production = journées produites ÷ jours travaillés. 100 % = une journée normale chaque jour travaillé, quel que soit le marché.", False),
    ("Les euros ne disparaissent pas : le CA est toujours calculé, par technicien et par marché, pour toi.", False),
    ("", False),
    ("CE QUE TU AS À FAIRE", True),
    ("1. Onglet « Règles » : vérifie les nombres des cases jaunes (combien dans une journée normale). FREE est à compléter.", False),
    ("2. Onglet « Saisies » : dans le Google Sheet, copie l'onglet Saisies de la ligne 2 jusqu'à la fin (colonnes A à N) et colle en A2. Tu peux aussi taper une ligne : date, technicien, projet, quantité et code BPU (liste en colonne G).", False),
    ("3. Onglet « Production » : choisis les dates (du… au…). Jours travaillés, journées produites, taux, classement et CA se calculent seuls.", False),
    ("", False),
    ("Cases jaunes, chiffres en bleu = à remplir. Le reste est calculé : ne pas modifier.", False),
    ("La colonne V de « Saisies » signale les lignes à vérifier (code inconnu, prix sur devis manquant, technicien absent de la liste…).", False),
    ("", False),
    ("POUR CHANGER UNE RÈGLE", True),
    ("• Changer un nombre : onglet « Règles », colonne D. Tout se recalcule.", False),
    ("• Une prestation mal comptée : onglet « Prestations », change sa catégorie (colonne F).", False),
    ("• Ajouter un technicien : onglet « Production », colonne A (15 places).", False),
    (f"• Plus de {SA2 - SA1 + 1} lignes de saisie : sélectionne la dernière ligne de « Saisies », colonnes O à V, et tire-la vers le bas.", False),
    ("• Onglet « Exemple » : quatre techniciens, une journée normale chacun sur quatre marchés, pour voir la différence avec les euros.", False),
]
for i, (t, gras) in enumerate(texte, start=3):
    c = ws.cell(i, 1, t); c.font = SOUS if gras else F(); c.alignment = Alignment(wrap_text=True, vertical='top')
ws.column_dimensions['A'].width = 130

# ---------- Production ----------
wpr = wb.create_sheet('Production')
wpr['A1'] = 'PRODUCTION DES TECHNICIENS'; wpr['A1'].font = TITRE
wpr['A3'] = 'Du'; wpr['A4'] = 'Au'
wpr['B3'] = datetime(2026, 10, 1)
wpr['B4'] = datetime(2026, 10, 31)
for a in ('A3', 'A4'): wpr[a].font = F(bold=True)
for a in ('B3', 'B4'):
    wpr[a].number_format = DATE; wpr[a].font = BLEU; wpr[a].fill = SAISIE; wpr[a].border = BORD
wpr['C3'] = '=COUNTIF(Saisies!$V$2:$V$100000,"?*")&" ligne(s) à vérifier dans l\'onglet Saisies (colonne V)"'
wpr['C3'].font = F(bold=True, color='C00000')
wpr['C4'] = "Choisis les dates : une semaine, un mois, l'année…"; wpr['C4'].font = NOTE
entetes(wpr, 6, ['Technicien', 'Jours travaillés', 'Journées produites', 'Taux de production',
                 'Classement', 'CA (€)', 'Classement au CA (ancienne méthode)', 'Lignes à vérifier'], 42)
periode = 'Saisies!$A:$A,">="&$B$3,Saisies!$A:$A,"<="&$B$4'
for k, r in enumerate(range(TE1, TE2 + 1)):
    nom = data['techs'][k] if k < len(data['techs']) else None
    f = {
        'A': nom,
        'B': f'=IF($A{r}="","",SUMIFS(Saisies!$U:$U,Saisies!$B:$B,$A{r},{periode}))',
        'C': f'=IF($A{r}="","",SUMIFS(Saisies!$S:$S,Saisies!$B:$B,$A{r},{periode}))',
        'D': f'=IF(OR($A{r}="",N(B{r})=0),"",C{r}/B{r})',
        'E': f'=IF(D{r}="","",RANK(D{r},$D${TE1}:$D${TE2}))',
        'F': f'=IF($A{r}="","",SUMIFS(Saisies!$R:$R,Saisies!$B:$B,$A{r},{periode}))',
        'G': f'=IF(OR($A{r}="",N(F{r})=0),"",RANK(F{r},$F${TE1}:$F${TE2}))',
        'H': f'=IF($A{r}="","",COUNTIFS(Saisies!$B:$B,$A{r},{periode},Saisies!$V:$V,"?*"))',
    }
    for col, v in f.items():
        c = wpr[f'{col}{r}']; c.value = v; c.font = F(); c.border = BORD
    wpr[f'A{r}'].font = BLEU; wpr[f'A{r}'].fill = SAISIE
    wpr[f'B{r}'].number_format = '0'; wpr[f'C{r}'].number_format = JOURS
    wpr[f'D{r}'].number_format = '0%'; wpr[f'F{r}'].number_format = EUROS
    for col in 'EGH': wpr[f'{col}{r}'].alignment = Alignment(horizontal='center')
rt = TE2 + 1
tot = {'A': 'Total équipe', 'B': f'=SUM(B{TE1}:B{TE2})', 'C': f'=SUM(C{TE1}:C{TE2})',
       'D': f'=IF(N(B{rt})=0,"",C{rt}/B{rt})', 'F': f'=SUM(F{TE1}:F{TE2})', 'H': f'=SUM(H{TE1}:H{TE2})'}
for col in 'ABCDEFGH':
    c = wpr[f'{col}{rt}']; c.value = tot.get(col); c.font = F(bold=True); c.fill = GRIS; c.border = BORD
wpr[f'B{rt}'].number_format = '0'; wpr[f'C{rt}'].number_format = JOURS
wpr[f'D{rt}'].number_format = '0%'; wpr[f'F{rt}'].number_format = EUROS
wpr.conditional_formatting.add(f'D{TE1}:D{TE2}', DataBarRule(start_type='num', start_value=0, end_type='num', end_value=1.5, color='5B9BD5'))
notes = [
    "Taux 100 % = une journée normale par jour travaillé, quel que soit le marché. Les congés et absences ne comptent pas comme jours travaillés.",
    "« Classement au CA » = l'ancienne façon de compter, en euros, laissée pour comparer.",
]
for i, t in enumerate(notes, start=rt + 1): wpr.cell(i, 1, t).font = NOTE

m0 = rt + 5
wpr.cell(m0 - 1, 1, 'PAR MARCHÉ, SUR LES MÊMES DATES').font = SOUS
entetes(wpr, m0, ['Marché', 'Journées produites', 'CA (€)', 'CA pour une journée de travail (€)'], 30)
for k, m in enumerate(marches):
    r = m0 + 1 + k
    f = {'A': m,
         'B': f'=SUMIFS(Saisies!$S:$S,Saisies!$C:$C,$A{r},{periode})',
         'C': f'=SUMIFS(Saisies!$R:$R,Saisies!$C:$C,$A{r},{periode})',
         'D': f'=IF(N(B{r})=0,"",C{r}/B{r})'}
    for col, v in f.items():
        c = wpr[f'{col}{r}']; c.value = v; c.font = F(); c.border = BORD
    wpr[f'B{r}'].number_format = JOURS; wpr[f'C{r}'].number_format = EUROS; wpr[f'D{r}'].number_format = EUROS
mfin = m0 + len(marches)
wpr.cell(mfin + 1, 1, "La dernière colonne montre ce que rapporte une journée de travail sur chaque marché : utile pour savoir quels marchés valent le coup.").font = NOTE
largeurs(wpr, {'A': 22, 'B': 13, 'C': 13, 'D': 16, 'E': 12, 'F': 13, 'G': 17, 'H': 11})
wpr.freeze_panes = 'B7'

# ---------- Saisies ----------
wsa = wb.create_sheet('Saisies')
cols_google = ['Date', 'Technicien', 'Projet', 'Prestation (si le projet en a plusieurs)', 'Quantité (vide = 1)',
               'Désignation', 'Code BPU', 'Unité', 'PU (€) (vide = prix du BPU)', 'Montant (€)',
               'N° ticket / référence', 'Commentaire', 'ID', 'Saisi le']
cols_calc = ['Catégorie (calculé)', 'Comment ça compte (calculé)', 'PU retenu (€) (calculé)', 'Montant (€) (calculé)',
             'Journées produites (calculé)', 'Travail (calculé)', 'Jour travaillé (calculé)', 'À vérifier (calculé)']
entetes(wsa, 1, cols_google + cols_calc, 45)
for j in range(len(cols_google) + 1, len(cols_google) + len(cols_calc) + 1):
    wsa.cell(1, j).fill = PatternFill('solid', fgColor='595959')
for r in range(SA1, SA2 + 1):
    im = f'MATCH($O{r},{R_A},0)'
    qte = f'IF($E{r}="",1,$E{r})'
    prix_bpu = f'INDEX({P_E},MATCH($G{r},{P_B},0))'
    f = {
        'O': f'=IF($G{r}="","",IFERROR(INDEX({P_F},MATCH($G{r},{P_B},0)),"?"))',
        'P': f'=IF($O{r}="","",IF($O{r}="?","?",IFERROR(INDEX({R_C},{im}),"?")))',
        'Q': f'=IF($G{r}="","",IF(ISNUMBER($I{r}),$I{r},IFERROR(IF(ISNUMBER({prix_bpu}),{prix_bpu},""),"")))',
        'R': f'=IFERROR(IF(ISNUMBER($Q{r}),$Q{r}*{qte},""),"")',
        'S': (f'=IFERROR(IF($P{r}="au nombre",IF(N(INDEX({R_D},{im}))>0,{qte}/INDEX({R_D},{im}),""),'
              f'IF($P{r}="au prix",IF(AND(ISNUMBER($R{r}),N(INDEX({R_E},{im}))>0),$R{r}/INDEX({R_E},{im}),""),'
              f'IF(OR($P{r}="ne compte pas",$P{r}="absence"),0,""))),"")'),
        'T': f'=IF(AND(ISNUMBER($A{r}),$B{r}<>"",$G{r}<>"",$P{r}<>"absence"),1,0)',
        'U': f'=IF($T{r}=1,IF(COUNTIFS($B${SA1}:$B{r},$B{r},$A${SA1}:$A{r},$A{r},$T${SA1}:$T{r},1)=1,1,0),0)',
        'V': (f'=IF(AND($A{r}="",$B{r}="",$G{r}=""),"",'
              f'IF(NOT(ISNUMBER($A{r})),"Date manquante ou non reconnue",'
              f'IF($B{r}="","Technicien manquant",'
              f'IF(COUNTIF(Production!$A${TE1}:$A${TE2},$B{r})=0,"Technicien absent de la liste (onglet Production)",'
              f'IF($G{r}="","Code BPU manquant (colonne G)",'
              f'IF($O{r}="?","Code BPU inconnu (onglet Prestations)",'
              f'IF($P{r}="?","Catégorie inconnue (onglet Règles)",'
              f'IF(AND($E{r}<>"",NOT(ISNUMBER($E{r}))),"Quantité non reconnue",'
              f'IF(AND($P{r}="au prix",NOT(ISNUMBER($Q{r}))),"Prix sur devis à saisir (colonne I)",'
              f'IF($S{r}="","Règle à compléter (onglet Règles)",""))))))))))'),
    }
    for col, v in f.items():
        c = wsa[f'{col}{r}']; c.value = v; c.font = F(); c.fill = GRIS
    wsa[f'A{r}'].number_format = DATE; wsa[f'N{r}'].number_format = 'dd/mm/yyyy hh:mm'
    for col in 'IJQ': wsa[f'{col}{r}'].number_format = EUROS2
    wsa[f'R{r}'].number_format = EUROS2; wsa[f'S{r}'].number_format = JOURS
    wsa[f'V{r}'].font = F(bold=True, color='C00000')
dv_t = DataValidation(type='list', formula1=f'=Production!$A${TE1}:$A${TE2}', allow_blank=True)
dv_c = DataValidation(type='list', formula1=f'={P_B}', allow_blank=True)
for dv, col in ((dv_t, 'B'), (dv_c, 'G')):
    dv.showErrorMessage = False; wsa.add_data_validation(dv); dv.add(f'{col}{SA1}:{col}{SA2}')
largeurs(wsa, {'A': 11, 'B': 18, 'C': 11, 'D': 22, 'E': 9, 'F': 34, 'G': 16, 'H': 12, 'I': 11, 'J': 11,
               'K': 14, 'L': 18, 'M': 10, 'N': 15, 'O': 24, 'P': 13, 'Q': 11, 'R': 11, 'S': 11, 'T': 8, 'U': 9, 'V': 34})
wsa.freeze_panes = 'C2'; wsa.auto_filter.ref = f'A1:V{SA2}'

# ---------- Règles ----------
wr = wb.create_sheet('Règles')
wr['A1'] = 'RÈGLES : COMBIEN DANS UNE JOURNÉE NORMALE ?'; wr['A1'].font = TITRE
wr['A2'] = "Change seulement les cases jaunes. « au nombre » : 1 unité = 1/(colonne D) de journée. « au prix » : prix de la prestation ÷ valeur d'une journée (colonne E)."
wr['A2'].font = NOTE
entetes(wr, 4, ['Catégorie', 'Ce qu\'on compte', 'Comment ça compte', 'Combien dans une journée normale',
                'Valeur d\'une journée (€)', 'Résultat', 'Remarque'], 42)
ligne_regle = {}
for i, (cat, unite, mode, n, rem) in enumerate(REGLES, start=RG1):
    ligne_regle[cat] = i
    for j, v in enumerate([cat, unite, mode, n, None, None, rem], 1):
        c = wr.cell(i, j, v); c.font = F(); c.border = BORD; c.alignment = WRAP
    wr.cell(i, 1).font = F(bold=True)
    if mode == 'au nombre':
        wr.cell(i, 4).fill = SAISIE; wr.cell(i, 4).font = BLEU
    wr.cell(i, 4).alignment = Alignment(horizontal='center', vertical='top')
    wr.cell(i, 5).number_format = EUROS
    wr.cell(i, 6).value = (f'=IF(C{i}="au nombre",IF(N(D{i})>0,"1 "&B{i}&" = 1/"&D{i}&" de journée","À COMPLÉTER : combien par jour ?"),'
                           f'IF(C{i}="au prix",IF(N(E{i})>0,"son prix ÷ "&ROUND(E{i},0)&" €","À COMPLÉTER"),'
                           f'IF(C{i}="absence","jour non travaillé","0 : ne compte pas")))')
for cat, (cat_ticket, code_ticket) in REF_AU_PRIX.items():
    i = ligne_regle[cat]
    wr.cell(i, 5).value = f'=IFERROR(D{ligne_regle[cat_ticket]}*INDEX({P_E},MATCH("{code_ticket}",{P_B},0)),"")'
largeurs(wr, {'A': 28, 'B': 13, 'C': 14, 'D': 13, 'E': 13, 'F': 30, 'G': 80})
wr.freeze_panes = 'B5'

# ---------- Prestations ----------
wp = wb.create_sheet('Prestations')
wp['A1'] = 'PRESTATIONS DU BPU ET CE QU\'ELLES VALENT EN JOURNÉES'; wp['A1'].font = TITRE
wp['A2'] = "Si une prestation est mal comptée, change sa catégorie en colonne F (cases jaunes). Les prix servent au CA et aux règles « au prix »."
wp['A2'].font = NOTE
entetes(wp, 4, ['Marché', 'Code BPU', 'Prestation', 'Unité', 'Prix client (€)', 'Catégorie', 'Vaut (journée) pour 1'], 30)
for i, (p, code, lib, unite, prix, cat) in enumerate(lignes, start=PR1):
    im = f'MATCH($F{i},{R_A},0)'
    part = (f'=IFERROR(IF(INDEX({R_C},{im})="au nombre",IF(N(INDEX({R_D},{im}))>0,1/INDEX({R_D},{im}),""),'
            f'IF(INDEX({R_C},{im})="au prix",IF(AND(ISNUMBER($E{i}),N(INDEX({R_E},{im}))>0),$E{i}/INDEX({R_E},{im}),""),0)),"")')
    for j, v in enumerate([p, code, lib, unite, prix, cat, part], 1):
        c = wp.cell(i, j, v); c.font = F(); c.border = BORD
    wp.cell(i, 5).number_format = EUROS2; wp.cell(i, 7).number_format = '0.00'
    wp.cell(i, 6).fill = SAISIE; wp.cell(i, 6).font = BLEU
dv = DataValidation(type='list', formula1=f'={R_A}', allow_blank=False)
dv.error = "Choisis une catégorie de l'onglet Règles."; wp.add_data_validation(dv); dv.add(f'F{PR1}:F{PR2}')
largeurs(wp, {'A': 11, 'B': 16, 'C': 60, 'D': 22, 'E': 12, 'F': 28, 'G': 12})
wp.freeze_panes = 'C5'; wp.auto_filter.ref = f'A4:G{PR2}'

# ---------- Exemple ----------
we = wb.create_sheet('Exemple')
we['A1'] = 'EXEMPLE : QUATRE TECHNICIENS, UNE JOURNÉE NORMALE CHACUN'; we['A1'].font = TITRE
we['A2'] = "Change le code (colonne C) ou la quantité (colonne E) pour tester. Les prix et les journées viennent des onglets Prestations et Règles."
we['A2'].font = NOTE
entetes(we, 4, ['Technicien', 'Marché', 'Code BPU', 'Ce qu\'il a fait', 'Quantité', 'Ce que paie le client (€)',
                'Journées produites', 'Classement en euros', 'Classement en journées'], 42)
exemple = [('Technicien A', 'LOSANGE-TICKET', 4), ('Technicien B', 'COVAGE-TICKET', 2),
           ('Technicien C', 'MC-T2', 7), ('Technicien D', 'CF-N1-STD', 4)]
e1, e2 = 5, 4 + len(exemple)
for i, (t, code, q) in enumerate(exemple, start=e1):
    trouve = f'MATCH($C{i},{P_B},0)'
    vals = [t, f'=IFERROR(INDEX(\'Prestations\'!$A${PR1}:$A${PR2},{trouve}),"Code inconnu")', code,
            f'=IFERROR(INDEX({P_C},{trouve}),"")', q,
            f'=IFERROR($E{i}*INDEX({P_E},{trouve}),"")', f'=IFERROR($E{i}*INDEX({P_G},{trouve}),"")',
            f'=IF(ISNUMBER(F{i}),RANK(F{i},$F${e1}:$F${e2}),"")', f'=IF(ISNUMBER(G{i}),RANK(G{i},$G${e1}:$G${e2}),"")']
    for j, v in enumerate(vals, 1):
        c = we.cell(i, j, v); c.font = F(); c.border = BORD
    for j in (3, 5): we.cell(i, j).font = BLEU; we.cell(i, j).fill = SAISIE
    we.cell(i, 6).number_format = EUROS; we.cell(i, 7).number_format = JOURS
    for j in (8, 9): we.cell(i, j).alignment = Alignment(horizontal='center')
for k, t in enumerate([
        "En euros, les quatre techniciens paraissent très différents : celui qui fait des heures AXIONE passe pour le moins bon.",
        "En journées, chacun a fait exactement une journée normale (1,00) : ils sont à égalité, ce qui est juste.",
        "C'est le marché qui rapporte plus ou moins, pas le technicien : ça se voit dans « Production », tableau « Par marché »."], start=e2 + 2):
    we.cell(k, 1, t).font = NOTE
largeurs(we, {'A': 14, 'B': 11, 'C': 16, 'D': 55, 'E': 10, 'F': 14, 'G': 12, 'H': 12, 'I': 12})

for s in wb.worksheets: s.sheet_view.showGridLines = False
wb.save(OUT)
print(len(lignes), 'prestations,', len(REGLES), 'règles ->', OUT)
