# Construit Devis_BPU_Axione_RHD57.xlsx au format du devis Circet (modèle Exemple_devis_LOSANGE.xlsx) :
# on garde l'en-tête, le tableau, les logos et les signatures du modèle, et on y met les articles et
# les prix des deux BPU Moselle Fibre RHD57 (Maintenance et VDR Optique) pour le client Axione.
#
#   python3 build_devis.py <BPU Maintenance.xlsx> <BPU VDR Optique.xlsx> <modèle devis.xlsx> <sortie.xlsx>
import math, re, sys
from copy import copy
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

SRC_MAINT, SRC_OPT, MODELE, OUT = sys.argv[1:5]

ORANGE = "FF6600"  # bandeau « lot » du BPU Maintenance
L1, LN = 21, 32    # lignes d'articles du devis modèle

# Unités du BPU raccourcies comme dans le modèle (Forf., ML, U…) pour tenir dans la colonne « U ».
UNITES = {"u": "U", "U": "U", "ml": "ML", "Forfait": "Forf.", "Forfait par opération": "Forf.",
          "Forfait par intervention": "Forf.", "Forfait par heure": "Heure", "Forfait global Mensuel": "Mois",
          "Par commande": "Cde", "Par mètre (hauteur nacelle)": "m (haut.)", "u (à la fibre)": "Fibre",
          "coef. appliqué au forfait": "Coef."}


def propre(t):
    return re.sub(r"[ \t]+", " ", str(t or "")).strip()


def lignes(t):
    return [propre(l) for l in str(t or "").split("\n") if propre(l)]


def prix(v):
    # Les BPU écrivent « sur devis », « Devis », « devis BPU »… : une seule écriture dans le devis.
    if v is None or isinstance(v, (int, float)):
        return v
    return "Sur devis"


# ---------- Lecture des BPU ----------
articles = []      # dict par prestation proposée dans le devis
majorations = []   # coefficients MC-07 / MC-08

ws = openpyxl.load_workbook(SRC_MAINT).worksheets[0]
lot = famille = ""
for r in range(2, ws.max_row + 1):
    a, code = ws.cell(r, 1).value, ws.cell(r, 2).value
    if a and not code:
        if ws.cell(r, 1).fill.fgColor.rgb == "FF" + ORANGE:
            lot = propre(a)
        else:
            famille = propre(a)
        continue
    if not code:
        continue
    d = dict(bpu="Maintenance", lot=lot, famille=famille, code=propre(code),
             desig=propre(ws.cell(r, 3).value), descr=" / ".join(lignes(ws.cell(r, 7).value)),
             unite=propre(ws.cell(r, 5).value), pbpu=prix(ws.cell(r, 6).value),
             comm=propre(ws.cell(r, 4).value), src=f"Maintenance BPU ST · BPU ligne {r}")
    if d["unite"].startswith("coef"):
        majorations.append(d)
        continue
    articles.append(d)
    if d["code"] == "RA-ON" and "0,93€/ml" in str(ws.cell(r, 7).value).replace(" ", ""):
        # Le descriptif de RA-ON prévoit 0,93 €/ml au-delà de 300 ml : on en fait une ligne.
        articles.append(dict(d, desig=d["desig"] + " – ml au-delà de 300 ml", unite="ml",
                             pbpu=0.93, descr="Mètre linéaire au-delà des 300 ml compris dans le forfait",
                             comm="D'après le descriptif de RA-ON (« jusqu'à 300ml, 0,93€/ml au-delà »)"))

ws = openpyxl.load_workbook(SRC_OPT).worksheets[0]
lot = famille = commun = ""
for r in range(4, ws.max_row + 1):
    a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
    if isinstance(a, str) and " - " in a:
        if re.match(r"^\d+ - ", a.strip()):
            lot, famille = propre(a), ""
        else:
            famille = propre(a)
        commun = ""
        continue
    if a is None:
        if b:
            commun = " ".join(lignes(b))  # descriptif commun à la famille
        continue
    code = str(a).strip()
    l = lignes(b)
    desig, reste = l[0].rstrip(" :"), l[1:]
    if reste and reste[0].startswith("Prix pour"):
        desig, reste = f"{desig} – {reste[0]}", reste[1:]
    unite = propre(ws.cell(r, 3).value) or "U"
    if unite == "fofait":  # coquille du BPU (4.6.1, 4.6.2)
        unite = "Forfait"
    if propre(ws.cell(r, 4).value) == "à la fibre":  # mesures et raccordements comptés à la fibre
        unite = f"{unite} (à la fibre)"
    articles.append(dict(bpu="VDR Optique", lot=lot, famille=famille or lot, code=code, desig=desig,
                         descr=" ".join(reste) or commun, unite=unite, pbpu=prix(ws.cell(r, 5).value),
                         comm=propre(ws.cell(r, 4).value),
                         src=f"BPU ST VDR Optique · BPU MISSION 3 Optique ligne {r}"))


def libelle(d, n=95):
    t = d["desig"]
    if len(t) > n:
        t = t[:n].rsplit(" ", 1)[0].rstrip(" ,;-–") + "…"
    return f"{d['code']} · {t}"


for d in articles:
    d["lib"] = libelle(d)
    assert d["pbpu"] is not None, d["code"]
libs = [d["lib"] for d in articles]
assert len(libs) == len(set(libs)), "libellés en double"

assert [m["code"] for m in majorations] == ["MC-07", "MC-08"]
MAJ_LIB = {"MC-07": "MC-07 · Majoration par technicien supplémentaire – sur la ligne au-dessus",
           "MC-08": "MC-08 · Majoration intervention en HNO – sur la ligne au-dessus"}
for m in majorations:
    m["lib"] = MAJ_LIB[m["code"]]
catalogue = articles + majorations
for d in catalogue:
    assert d["unite"] in UNITES, d["unite"]

wb = openpyxl.load_workbook(MODELE)
dv, bpu, aide = wb["Devis"], wb["BPU 2026"], wb["Mode d'emploi"]

# ---------- Onglet BPU : même présentation que « BPU 2026 », articles Axione ----------
bpu.title = "BPU RHD57"
FIN = 3 + len(catalogue)
bpu["A1"] = "Bordereau de Prix Unitaires — Moselle Fibre RHD57 — Axione (Maintenance et VDR Optique)"
bpu["A2"] = ("Catalogue de référence — sert de source aux menus déroulants de l'onglet \"Devis\". "
             "Ne pas modifier l'ordre des colonnes B, C et D.")
ent = {c: copy(bpu[f"{c}3"]._style) for c in "ABCD"}
cel = {c: copy(bpu[f"{c}4"]._style) for c in "ABCD"}
bpu.delete_rows(4, bpu.max_row)
for c, t, w in [("A", "Catégorie", 48), ("B", "Article", 95), ("C", "Unité", 10), ("D", "PU HT", 14),
                ("E", "Unité BPU", 26), ("F", "Descriptif", 70), ("G", "Source", 46)]:
    bpu[f"{c}3"] = t
    bpu[f"{c}3"]._style = copy(ent.get(c, ent["B"]))
    bpu.column_dimensions[c].width = w
for i, d in enumerate(catalogue, 4):
    vals = {"A": f"{d['bpu']} – {d['famille']}", "B": d["lib"], "C": UNITES[d["unite"]], "D": d["pbpu"],
            "E": d["unite"], "F": " / ".join(x for x in [d["desig"] if d in majorations else "", d["descr"]] if x),
            "G": d["src"]}
    for c, v in vals.items():
        bpu[f"{c}{i}"] = v
        bpu[f"{c}{i}"]._style = copy(cel.get(c, cel["B"]))
    if d in majorations:
        bpu[f"D{i}"].number_format = "0.0000"
bpu.auto_filter.ref = f"A3:G{FIN}"
bpu.freeze_panes = "A4"

# ---------- Onglet Devis : en-tête Axione ----------
dv["B8"] = "Bordereau de prix unitaire Moselle Fibre – RHD57"
dv["B11"] = "Département : 57"
dv["B12"] = "Objet : Intervention suite au ticket "
dv["B13"] = "Devis n° : DEVAX57-"
dv["B14"] = "Affaire suivie par : C LECOMTE"
dv["B15"] = "A l'attention de : Nicolas FRUTOS"
dv["B17"] = "Réf. ticket Axione : "
dv["B19"] = "PRESTATIONS SELON BPU MOSELLE FIBRE – RHD57 (Maintenance et VDR Optique)"
dv["H41"] = "Pour Axione"

# Le logo Losange laisse la place au nom du client (le logo Circet et la signature restent).
dv._images = [im for im in dv._images if im.anchor._from.col != 7]
dv.merge_cells("H2:K4")
dv["H2"] = "AXIONE"
dv["H2"].font = Font(name="Aptos Narrow", size=22, bold=True, color="1F497D")
dv["H2"].alignment = Alignment(horizontal="right", vertical="center")
dv.merge_cells("H5:K5")
dv["H5"] = "Moselle Fibre – RHD57"
dv["H5"].font = Font(name="Aptos Narrow", size=10, italic=True, color="808080")
dv["H5"].alignment = Alignment(horizontal="right", vertical="center")

# Lignes d'articles : unité et prix depuis le BPU. Un article « Sur devis » (ou hors BPU) laisse
# le P.U.H.T à taper. Une majoration MC-07 / MC-08 vaut P.U.H.T de la ligne au-dessus × (coef − 1).
plage = lambda c: f"'BPU RHD57'!${c}$4:${c}${FIN}"
for r in range(L1, LN + 1):
    m = f"MATCH($B{r},{plage('B')},0)"
    dv[f"H{r}"] = f'=IF($B{r}="","",IFERROR(INDEX({plage("C")},{m}),""))'
    dv[f"I{r}"] = (f'=IF($B{r}="","",IFERROR(IF(INDEX({plage("C")},{m})="Coef.",'
                   f'ROUND(I{r - 1}*(INDEX({plage("D")},{m})-1),2),ROUND(INDEX({plage("D")},{m}),2)),"Sur devis"))')
    dv[f"K{r}"] = f'=IF($B{r}="","",IF(ISNUMBER(I{r}),ROUND(I{r}*N(J{r}),2),"PU à saisir"))'
    dv[f"H{r}"]._style = copy(dv[f"H{L1 + 1}"]._style)
    dv[f"H{r}"].alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
v = DataValidation(type="list", formula1=f"={plage('B')}", allow_blank=True, errorStyle="warning",
                   errorTitle="Article hors BPU",
                   error="Cet article n'est pas dans le BPU. Continuer pour le saisir quand même, "
                         "puis taper le P.U.H.T.")
v.add(f"B{L1}:B{LN}")
dv.add_data_validation(v)
dv.conditional_formatting.add(f"I{L1}:I{LN}", FormulaRule(formula=[f'I{L1}="Sur devis"'],
                                                         fill=PatternFill("solid", fgColor="FFFF00")))
dv.conditional_formatting.add(f"K{L1}:K{LN}", FormulaRule(formula=[f'K{L1}="PU à saisir"'],
                                                         font=Font(color="C00000", bold=True)))

# ---------- Mode d'emploi ----------
st = {k: copy(aide[ref]._style) for k, ref in [("titre", "A1"), ("section", "A3"), ("texte", "A4")]}
aide.delete_rows(1, aide.max_row)
texte = [
    ("titre", "Comment utiliser ce classeur"),
    ("section", "1. Renseigner l'en-tête"),
    ("texte", "Dans l'onglet \"Devis\", complète les cases du haut : Département, Objet, Devis n°, Affaire suivie "
              "par, A l'attention de (Nicolas FRUTOS chez Axione), puis la référence du ticket dans la bande bleue."),
    ("section", "2. Ajouter les lignes de prestations"),
    ("texte", "Clique sur une cellule \"Article\" : un menu déroulant propose tous les articles du BPU Moselle "
              "Fibre RHD57, sous la forme « code · désignation » (Maintenance puis VDR Optique). Choisis l'article : "
              "l'unité et le P.U.H.T se remplissent tout seuls. Indique la quantité, le Montant HT se calcule."),
    ("texte", "Astuce : tape le début du code (ex. MC-T1 ou 4.3.2) pour retrouver l'article plus vite."),
    ("section", "3. Articles « Sur devis »"),
    ("texte", "Pour les articles sans prix dans le BPU (RA-RO, MC-T6, PU-01 à PU-06) ou hors BPU, le P.U.H.T "
              "affiche « Sur devis » sur fond jaune : tape ton prix directement dans la case."),
    ("section", "4. Majorations MC-07 et MC-08"),
    ("texte", "Choisis la ligne de majoration juste sous l'intervention concernée, avec la même quantité. Son P.U.H.T "
              "vaut le P.U.H.T de la ligne au-dessus × (coefficient − 1) : par exemple MC-08 (HNO, coef. 1,4078) "
              "ajoute 40,78 % au forfait. Les coefficients sont dans l'onglet \"BPU RHD57\"."),
    ("section", "5. Le total"),
    ("texte", "\"Sous Total HT\" additionne automatiquement toutes les lignes remplies."),
    ("section", "6. Pour un nouveau devis"),
    ("texte", "Clic droit sur l'onglet \"Devis\" → Déplacer ou copier → cocher \"Créer une copie\" → renomme "
              "l'onglet avec le n° de ticket. Ça garde une trace de chaque devis dans le même classeur, avec l'onglet "
              "\"BPU RHD57\" partagé par tous."),
    ("section", "7. Mettre à jour les prix"),
    ("texte", "Si le BPU change, modifie uniquement l'onglet \"BPU RHD57\" (colonnes C et D) — tous les devis "
              "qui utilisent un article existant se mettent à jour tout seuls."),
]
r = 1
for i, (k, t) in enumerate(texte):
    if k == "section" and i > 1:
        aide.row_dimensions[r].height = 9.75
        r += 1
    aide.cell(r, 1, t)._style = copy(st[k])
    aide.row_dimensions[r].height = 24 if k != "texte" else 15 * math.ceil(len(t) / 105) + 6
    r += 1
    if k == "titre":
        aide.row_dimensions[r].height = 9.75
        r += 1

wb.save(OUT)
print(f"{len(articles)} articles ({sum(d['bpu'] == 'Maintenance' for d in articles)} maintenance, "
      f"{sum(d['bpu'] == 'VDR Optique' for d in articles)} optique) + {len(majorations)} majorations -> {OUT}")
